"""Puente JSONL de Detalles: motor original aislado, sin Tk ni dependencias Qt."""
import argparse
from contextlib import contextmanager
import json
import logging
from pathlib import Path
import re
import sqlite3
import sys
import threading
import time

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / 'src'))

from core.catalog import TableCatalog, table_name
from core.captcha import load_captcha_lookup
from core.job_store import JobStore, RunLock, active_revisions
from core.messages import pending_reason
from core.playwright_processor import PlaywrightDUAProcessor
from flows.detalles_flow import DetallesFlow
from daily import DailyCoordinator, local_now
import utils


class Protocol:
    def __init__(self):
        self.lock = threading.Lock()

    def emit(self, event, **data):
        with self.lock:
            sys.stdout.write(json.dumps({'event': event, **data}, ensure_ascii=False, allow_nan=False) + '\n')
            sys.stdout.flush()


class EventLog(logging.Handler):
    def __init__(self, protocol):
        super().__init__()
        self.protocol = protocol

    def emit(self, record):
        if not getattr(record, 'technical_only', False):
            self.protocol.emit('log', message=record.getMessage(), level=getattr(record, 'tag', record.levelname),
                               time=time.strftime('%H:%M:%S'))


class DetailsService:
    def __init__(self, storage_root, protocol, offline=False):
        self.root = Path(storage_root).resolve()
        self.protocol = protocol
        # Conserva estructura y formatos del bot; --storage-root aísla pruebas.
        utils._BASE_DIR = str(self.root)
        utils._CONFIG_DIR = str(self.root / 'config')
        utils._MACHINE_PATH = str(self.root / 'config' / 'machine.json')
        utils._SESSION_PATH = str(self.root / 'config' / 'session.json')
        self.machine = utils.load_machine_config()
        self.session = utils.load_session()
        self.catalog = TableCatalog(self.root / 'data' / 'queues')
        self.catalog.discover_legacy()
        self.fields = list(DetallesFlow().extract_data('').columns)
        self.captcha = load_captcha_lookup()
        self.stop_event = threading.Event()
        self.active_id = self.session.get('active_table_id', '')
        self.daily = DailyCoordinator(self, offline=offline)
        if not self.active_id and self.session.get('input_file_path') and self.session.get('sheet_name'):
            self.active_id = utils.dataset_key(self.session['input_file_path'], self.session['sheet_name'])
        try:
            self.catalog.get(self.active_id)
        except (ValueError, OSError):
            self.active_id = ''
        utils.setup_logging(str(self.root / 'logs' / 'workbench.log'))
        logging.getLogger().addHandler(EventLog(protocol))

    @contextmanager
    def read_database(self, record):
        db = sqlite3.connect(self.catalog.database(record).as_uri() + '?mode=ro', uri=True, timeout=2)
        db.row_factory = sqlite3.Row
        try:
            yield db
        finally:
            db.close()

    def summary(self, record):
        if not record:
            return {state: 0 for state in ('pending', 'running', 'retry', 'succeeded', 'invalid', 'total', 'attempts', 'source_rows', 'retained')}
        # Instantánea de solo lectura: nunca recupera ni modifica reservas ajenas.
        with self.read_database(record) as db:
            revisions = active_revisions(db)
            marks = ','.join('?' for _ in revisions)
            counts = dict(db.execute(f'''SELECT status, COUNT(*) FROM jobs WHERE status!='invalid'
                OR EXISTS (SELECT 1 FROM source_rows s WHERE s.job_id=jobs.id AND s.revision IN ({marks}))
                GROUP BY status''', revisions).fetchall())
            source = db.execute("SELECT value FROM metadata WHERE key='source_rows'").fetchone()
            attempts = db.execute('SELECT COALESCE(SUM(attempts),0) FROM jobs').fetchone()[0]
            retained = db.execute('SELECT COUNT(*) FROM retained_results').fetchone()[0]
        return {**{state: counts.get(state, 0) for state in ('pending', 'running', 'retry', 'succeeded', 'invalid')},
                'total': sum(counts.values()), 'attempts': attempts,
                'source_rows': json.loads(source[0]) if source else 0, 'retained': retained}

    def record(self):
        return self.catalog.get(self.active_id) if self.active_id else None

    def paths(self, record=None):
        record = record or self.record()
        base = self.machine.get('output_base_path', '').strip()
        if not base:
            raise ValueError('Selecciona una carpeta de resultados antes de iniciar o exportar.')
        if not record:
            raise ValueError('Importa o abre una tabla guardada.')
        name = self.output_name(record['name'])
        for other in self.catalog.list_tables():
            if other['id'] != record['id'] and self.output_name(other['name']).casefold() == name.casefold():
                raise ValueError('Dos tablas producen el mismo nombre de carpeta. Renombra una antes de exportar.')
        folder = Path(base).resolve() / name
        internal = self.root / 'data' / 'exports' / record['id']
        folder.mkdir(parents=True, exist_ok=True)
        internal.mkdir(parents=True, exist_ok=True)
        return {
            'output_csv': str(folder / f'Detalles_DUAs_{name}.csv'),
            'processed_log': str(internal / 'procesados.txt'),
            'failed_duas': str(internal / 'pendientes.txt'),
            'pending_json': str(internal / 'pendientes.json'),
            'diagnostics': str(internal / 'diagnosticos'),
        }

    @staticmethod
    def output_name(name):
        # Preserve the chosen name, including spaces and accents. Replace only
        # characters Windows cannot use in a directory or file name.
        safe = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', name).strip(' .') or 'Tabla'
        if safe.split('.')[0].upper() in {'CON', 'PRN', 'AUX', 'NUL', *{f'COM{i}' for i in range(1, 10)}, *{f'LPT{i}' for i in range(1, 10)}}:
            safe = '_' + safe
        return safe

    def select(self, table_id):
        record = self.catalog.get(table_id)
        self.active_id = record['id']
        self.session.update(active_table_id=record['id'], input_file_path=record['source_path'], sheet_name=record['source_sheet'])
        utils.save_session(self.session)

    def snapshot(self):
        try:
            record = self.record()
        except (ValueError, OSError):
            self.active_id = ''
            self.session['active_table_id'] = ''
            utils.save_session(self.session)
            record = None
        tables = self.catalog.list_tables()
        for item in tables:
            try:
                item['summary'] = self.summary(item)
            except (OSError, sqlite3.Error):
                item['summary'] = self.summary(None)
                item['unavailable'] = True
        return {'tables': tables, 'activeTable': record or {}, 'summary': self.summary(record),
                'config': self.machine, 'session': self.session, 'fields': self.fields, 'daily': self.daily.snapshot(),
                'storageRoot': str(self.root),
                'outputFolder': str(Path(self.paths(record)['output_csv']).parent) if record and self.machine.get('output_base_path') else ''}

    def runtime(self):
        from playwright.sync_api import sync_playwright
        with sync_playwright() as playwright:
            browser = Path(playwright.chromium.executable_path)
        return {'chromiumReady': browser.is_file(), 'chromiumPath': str(browser), 'captchaEntries': len(self.captcha)}

    def configure(self, values):
        updated = dict(self.machine)
        for key, low, high in [('cant_hilos', 1, 20), ('timeout', 5, 120), ('max_retries', 1, 10)]:
            value = values.get(key, updated.get(key, {'cant_hilos': 1, 'timeout': 30, 'max_retries': 3}[key]))
            if isinstance(value, bool) or not isinstance(value, int) or not low <= value <= high:
                raise ValueError(f'Valor inválido para {key}: debe estar entre {low} y {high}.')
            updated[key] = value
        for key in ('headless', 'dark_mode', 'reuse_context', 'lightweight_details', 'daily_mode'):
            if key in values:
                if not isinstance(values[key], bool):
                    raise ValueError(f'Valor inválido para {key}.')
                updated[key] = values[key]
        if 'output_base_path' in values:
            updated['output_base_path'] = str(values['output_base_path']).strip()
        if 'daily_start_date' in values:
            from datetime import date
            start = date.fromisoformat(str(values['daily_start_date']))
            if start >= local_now().date():
                raise ValueError('La fecha inicial debe ser anterior a hoy.')
            updated['daily_start_date'] = start.isoformat()
        utils.save_machine_config(updated)
        self.machine = updated

    def view(self, request):
        record = self.record()
        if not record:
            return {'rows': [], 'total': 0, 'offset': 0}
        offset = max(0, int(request.get('offset', 0)))
        total, rows = self.catalog.read_rows(record['id'], request.get('filter', 'all'), str(request.get('search', '')),
                                             offset=offset, limit=100)
        for row in rows:
            row['reason'] = ('Resultado guardado.' if row['status'] == 'succeeded' else
                             'Consulta en curso.' if row['status'] == 'running' else
                             'Listo para consultar.' if row['status'] == 'pending' and not row['last_error'] else
                             pending_reason(row['error_kind'], row['last_error']))
        return {'rows': rows, 'total': total, 'offset': offset}

    def detail(self, job_id):
        record = self.record()
        if not record:
            raise ValueError('No hay tabla activa.')
        with self.read_database(record) as db:
            row = db.execute('SELECT * FROM jobs WHERE id=?', (job_id,)).fetchone()
            if row is None:
                raise ValueError('La fila ya no existe.')
            sources = [dict(item) for item in db.execute('SELECT revision,row_number,raw FROM source_rows WHERE job_id=?', (job_id,))]
            attempts = [dict(item) for item in db.execute('SELECT started,finished,outcome,error FROM attempts WHERE job_id=? ORDER BY id DESC LIMIT 100', (job_id,))]
        result = json.loads(row['result']) if row['result'] else {}
        evidence = row['evidence']
        if evidence and not Path(evidence).is_file():
            archived = self.root / 'data' / 'exports' / record['id'] / 'legacy-output' / 'diagnosticos' / Path(evidence).name
            if archived.is_file():
                evidence = str(archived)
        return {**dict(row), 'result': result, 'sources': sources, 'attemptHistory': attempts,
                'evidence': evidence,
                'values': [{'field': name, 'value': str(result.get(name, ''))} for name in self.fields],
                'reason': 'Resultado guardado.' if result else 'Consulta en curso.' if row['status'] == 'running' else
                          pending_reason(row['error_kind'], row['last_error']) if row['last_error'] else 'Listo para consultar.'}

    def execute(self, command, request):
        if command == 'daily_calendar':
            month = str(request.get('month') or self.daily.month)
            if not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', month) or not 1 <= int(month[:4]) <= 9999:
                raise ValueError('Mes inválido.')
            self.daily.month = month
            return {'daily': self.daily.snapshot()}
        if command == 'daily_run':
            return self.daily.run()
        if command == 'initialize':
            return {'snapshot': self.snapshot(), 'runtime': self.runtime(), 'view': self.view({})}
        if command == 'inspect_excel':
            from openpyxl import load_workbook
            source = Path(request['path']).resolve()
            book = load_workbook(source, read_only=True)
            try:
                sheets = book.sheetnames
            finally:
                book.close()
            return {'fileInfo': {'path': str(source), 'sheets': sheets}}
        if command == 'prepare_import':
            name = table_name(request['name'])
            target = self.catalog.find_by_name(name)
            return {'importPlan': {'name': name, 'targetId': target['id'] if target else None,
                                   'source': request['source'], 'sheet': request['sheet'], 'appended': bool(target)}}
        if command == 'import':
            if 'targetId' not in request:
                raise ValueError('Confirma la tabla de destino antes de importar.')
            source, sheet, name = request['source'], request['sheet'], request['name']
            base = self.machine.get('output_base_path', '')
            legacy = Path(base) / sheet / f'MM20YY_Detalles_DUAs__procesados_{sheet}.csv' if base else None
            record = self.catalog.import_excel(source, sheet, name, legacy, self.fields, expected_target=request['targetId'])
            self.select(record['id'])
            report = record['import_summary']
            logging.info('Tabla «%s»: %s DUAs nuevos, %s existentes y %s entradas para revisar.', name, report['added'], report['existing'], report['invalid'])
            return {'snapshot': self.snapshot(), 'view': self.view(request.get('viewQuery', {})), 'importSummary': report}
        if command == 'select':
            self.select(request['tableId'])
        elif command == 'configure':
            self.configure(request['config'])
        elif command == 'view':
            return {'view': self.view(request)}
        elif command == 'detail':
            return {'detail': self.detail(request['jobId'])}
        elif command == 'rename':
            self.catalog.rename(self.active_id, request['name'])
        elif command in {'remove_rows', 'clean', 'delete'}:
            if request.get('confirmTableId') != self.active_id or not self.active_id:
                raise ValueError('La tabla seleccionada cambió. Confirma de nuevo la operación.')
            if command == 'delete':
                backup = self.catalog.delete_table(self.active_id)
                self.active_id = ''
                self.session['active_table_id'] = ''
                utils.save_session(self.session)
            else:
                count, backup = self.catalog.remove_rows(self.active_id, ids=request.get('ids'), succeeded_only=command == 'clean')
                logging.info('Se retiraron %s filas. Los resultados confirmados se conservan.', count)
            logging.info('Respaldo previo: %s', backup)
        elif command == 'export':
            record = self.record()
            if not record:
                raise ValueError('Selecciona una tabla.')
            with RunLock(self.catalog.database(record)):
                store = JobStore(self.catalog.database(record))
                try:
                    store.export(self.paths(record), self.fields)
                finally:
                    store.close()
            logging.info('Resultados exportados desde SQLite; no se consultó el portal.')
        elif command == 'run':
            record = self.record()
            if not record:
                raise ValueError('Selecciona una tabla.')
            paths = self.paths(record)
            with RunLock(self.catalog.database(record)):
                self.catalog.get(record['id'])
                store = JobStore(self.catalog.database(record))
                monitor_stop = threading.Event()
                def monitor():
                    last = None
                    while not monitor_stop.wait(0.2):
                        value = store.summary()
                        if value != last:
                            self.protocol.emit('progress', summary=value)
                            last = value
                monitor_thread = threading.Thread(target=monitor, daemon=False)
                try:
                    processor = PlaywrightDUAProcessor(self.machine, paths, DetallesFlow(), self.captcha, self.stop_event)
                    monitor_thread.start()
                    processor.run(store, lambda summary: self.protocol.emit('progress', summary=summary))
                finally:
                    monitor_stop.set()
                    if monitor_thread.is_alive():
                        monitor_thread.join()
                    summary = store.summary()
                    store.close()
            return {'snapshot': self.snapshot(), 'view': self.view(request.get('viewQuery', {})),
                    'runStatus': 'stopped' if self.stop_event.is_set() else 'completed' if summary['total'] == summary['succeeded'] else 'pending'}
        elif command == 'refresh':
            self.machine = utils.load_machine_config()
        else:
            raise ValueError(f'Comando desconocido: {command}')
        return {'snapshot': self.snapshot(), 'view': self.view(request.get('viewQuery', {}))}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--storage-root', type=Path, default=ROOT)
    parser.add_argument('--offline-fixture', action='store_true', help='Solo pruebas: sustituir el portal por HTML local y bloquear red.')
    args = parser.parse_args()
    if args.offline_fixture:
        if args.storage_root.resolve() == ROOT:
            parser.error('La fixture exige almacenamiento separado de la instalación normal.')
        import asyncio
        async def local_result(processor, page, row):
            await page.route('**/*', lambda route: route.abort())
            await asyncio.sleep(0.25)
            html = ''.join(f'<span id="{key}">{value}</span>' for key, value in [
                ('span_CODI_ADUAN', row['Aduana']), ('span_ANO_PRESE', row['Año']), ('span_NUME_CORRE', row['Número'])])
            await page.set_content(html)
            return await page.content()
        PlaywrightDUAProcessor._fetch_result = local_result
    sys.stdout.reconfigure(encoding='utf-8', line_buffering=True)
    sys.stdin.reconfigure(encoding='utf-8')
    protocol = Protocol()
    try:
        service = DetailsService(args.storage_root, protocol, offline=args.offline_fixture)
    except Exception as error:
        protocol.emit('fatal', message=str(error))
        return 1
    operation = None
    busy = threading.Event()
    readers = []

    def execute(request, foreground=True):
        try:
            result = service.execute(request['command'], request)
            protocol.emit('result', requestId=request.get('requestId'), command=request['command'], **result)
        except Exception as error:
            logging.exception('Operación interrumpida: %s', error, extra={'technical_only': True})
            data = {}
            if foreground:
                try:
                    data = {'snapshot': service.snapshot(), 'view': service.view(request.get('viewQuery', {}))}
                except Exception:
                    pass
            protocol.emit('error', requestId=request.get('requestId'), command=request.get('command'), message=str(error), **data)
        finally:
            if foreground:
                busy.clear()
            protocol.emit('idle', requestId=request.get('requestId'), readonly=not foreground)

    for line in sys.stdin:
        try:
            request = json.loads(line)
            command = request['command']
            if command == 'stop':
                service.stop_event.set()
                protocol.emit('stopping')
                continue
            if command == 'shutdown':
                break
            if busy.is_set() and command in {'view', 'detail', 'daily_calendar'}:
                readers = [thread for thread in readers if thread.is_alive()]
                if len(readers) >= 2:
                    protocol.emit('error', requestId=request.get('requestId'), command=command, message='Espera a que termine la lectura actual.')
                    continue
                thread = threading.Thread(target=execute, args=(request, False), daemon=False)
                readers.append(thread)
                thread.start()
                continue
            if busy.is_set():
                protocol.emit('error', requestId=request.get('requestId'), command=command,
                              message='Espera a que termine la operación actual o detén las consultas.')
                continue
            service.stop_event.clear()
            busy.set()
            operation = threading.Thread(target=execute, args=(request,), daemon=False)
            operation.start()
        except Exception as error:
            protocol.emit('error', message=str(error))
    service.stop_event.set()
    if operation:
        operation.join()  # Espera cancelación, exportación y cierre ordenado.
    for reader in readers:
        reader.join()
    protocol.emit('closed')
    return 0


if __name__ == '__main__':
    sys.exit(main())
