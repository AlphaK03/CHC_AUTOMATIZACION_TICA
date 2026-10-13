"""Comprueba el generador real y BLs_BOT juntos, sin consultar TICA."""
import argparse
import asyncio
from contextlib import closing
from itertools import islice
import json
from pathlib import Path
import sqlite3
import subprocess
import sys
import tempfile
from threading import Event

from openpyxl import load_workbook, Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from core.engine import BLSProcessor
from core.inputs import batch_scope, input_job, query_dates
from input_cache import save_input_cache
from input_files import BLS_COLUMNS, input_sheets, load_input_file


GENERATE = '''
import csv, json, sys
from pathlib import Path
sys.path.insert(0, sys.argv[1])
from bls_base import build_base, Source
from conocimientos import COLUMNS
folder=Path(sys.argv[2])
rows=[
 ['4-0-26000125-P010-01/04/26','001','00001234567890123456789','N'],
 ['1-0-20250024-P027-02/02/26','0002','000001234','N'],
 ['1-0-20240739-P027-02/01/25','003','001-000023547860','N'],
 ['4-0-26000125-P010-01/04/26','001','00001234567890123456789','N'],
 ['4-0-26000125-P010-01/04/26','005','courier_excluido','S'],
 ['4-0-26000125-P010-01/04/26','004','conflicto_a','N'],
 ['4-0-26000125-P010-01/04/26','004','conflicto_b','N']]
source=folder/'conocimientos.csv'
with source.open('w',encoding='utf-8-sig',newline='') as stream:
 writer=csv.writer(stream);writer.writerow(COLUMNS)
 writer.writerows(row+['']*(len(COLUMNS)-4) for row in rows)
reports=[build_base([Source(source,'CSV')],folder/('base'+extension),audit_dir=folder/'auditoria')
         for extension in ('.xlsx','.sqlite3')]
assert all(r['exported']==3 and r['conflicts']==1 for r in reports), reports
print(json.dumps({'exported':3,'conflicts':1}))
'''


def write_sqlite(path, rows):
    with closing(sqlite3.connect(path)) as db:
        with db:
            db.execute('CREATE TABLE BLS (' + ','.join('"' + c + '" TEXT NOT NULL' for c in BLS_COLUMNS) + ')')
            db.executemany('INSERT INTO BLS VALUES (?,?,?,?,?)', rows)


async def check_resume(excel, sql, folder):
    config = {'sheet_name': 'BLS', 'output_base_path': str(folder), 'cant_hilos': 2,
              'max_retries': 1, 'dates_from_input': True}
    assert input_sheets(excel) == input_sheets(sql) == ['BLS']
    a, b = load_input_file(excel), load_input_file(sql)
    assert a.to_dict('records') == b.to_dict('records')
    assert batch_scope(a, config) == batch_scope(b, config)
    for ordinal, record in enumerate(a.to_dict('records'), 1):
        assert not input_job(record, ordinal)['error'], record
        query_dates(record, config)
    cache = folder / 'cache.sqlite3'
    save_input_cache(b, cache, sheet_name='BLS')
    assert batch_scope(load_input_file(cache), config) == batch_scope(a, config)
    calls = []
    class Session:
        async def query(self, record):
            calls.append((record['BL original'], record['Secuencia Tica']))
            return {'status': 'success', 'columns': ['BL', 'Aduana'],
                    'rows': [[record['BL original'], '013']]}
        async def close(self):
            pass
    first = await BLSProcessor(config, Event()).run_async(a, session_factory=lambda worker: Session())
    assert first['saved'] == len(a)
    assert len(calls) == len(a)
    second = await BLSProcessor(config, Event()).run_async(b, session_factory=lambda worker: Session())
    assert second['saved'] == len(a) and len(calls) == len(a)
    return {'rows': len(a), 'identical_inputs': True, 'valid_jobs': True,
            'same_batch': True, 'resume_without_repeat': True}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--manifiestos', type=Path, required=True)
    parser.add_argument('--sqlite-real', type=Path)
    parser.add_argument('--excel-real', type=Path)
    parser.add_argument('--informe', type=Path, required=True)
    args = parser.parse_args()
    if bool(args.sqlite_real) != bool(args.excel_real):
        parser.error('Indique ambas bases reales del mismo alcance o ninguna.')
    with tempfile.TemporaryDirectory(prefix='compatibilidad_bls_') as temporary:
        folder = Path(temporary)
        generated = subprocess.run([str(args.manifiestos / '.venv/Scripts/python.exe'), '-c', GENERATE,
                                   str(args.manifiestos / 'src'), str(folder)],
                                  capture_output=True, text=True, timeout=60)
        if generated.returncode:
            raise RuntimeError(generated.stderr or generated.stdout)
        report = {'generator': json.loads(generated.stdout),
                  'generated': asyncio.run(check_resume(folder/'base.xlsx', folder/'base.sqlite3', folder))}
        if args.sqlite_real:
            # Muestra acotada; las entradas reales permanecen de solo lectura.
            with closing(sqlite3.connect(args.sqlite_real.resolve().as_uri()+'?mode=ro', uri=True)) as db:
                db.execute('PRAGMA query_only=ON')
                rows = db.execute('SELECT * FROM BLS ORDER BY rowid LIMIT 12').fetchall()
            wb = load_workbook(args.excel_real, read_only=True, data_only=False)
            try:
                values = wb['BLS'].values
                assert list(next(values)) == BLS_COLUMNS
                assert list(islice(values, 12)) == rows, 'Las bases reales no tienen la misma muestra inicial.'
            finally:
                wb.close()
            real = folder/'real'
            real.mkdir()
            write_sqlite(real/'base.sqlite3', rows)
            book = Workbook()
            book.active.title = 'BLS'
            book.active.append(BLS_COLUMNS)
            for row in rows:
                book.active.append(row)
            book.save(real/'base.xlsx')
            book.close()
            report['real_sample'] = asyncio.run(check_resume(real/'base.xlsx', real/'base.sqlite3', real))
        report.update(ready=True, tica_queries=0, production_progress_changed=False)
    args.informe.parent.mkdir(parents=True, exist_ok=True)
    args.informe.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
