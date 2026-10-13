import csv
import hashlib
import io
import json
import os
import sqlite3
import subprocess
import sys
import tempfile
import time
import unittest
from contextlib import closing, redirect_stdout
from pathlib import Path
from threading import Event
from unittest.mock import patch

from openpyxl import load_workbook

from bls_base import (BuildCancelled, available_periods, build_base,
                      default_selection, discover_sources, inspect_sources)
from conocimientos import COLUMNS, GridPage, KnowledgeStore


JAN = '1-0-20240739-P027-02/01/25'
FEB = '1-0-20250024-P027-02/02/25'
AIR = '4-0-26000124-P010-01/04/26'


def record(mid, sequence='001', bl='0000001234567890123456789', courier='N'):
    return [mid, sequence, bl, courier] + [''] * (len(COLUMNS) - 4)


def write_csv(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream)
        writer.writerow(COLUMNS)
        writer.writerows(rows)
    return path


class DownloadedPeriodTests(unittest.TestCase):
    def test_mixed_file_periods_use_arrival_date_instead_of_filename_or_number(self):
        with tempfile.TemporaryDirectory() as folder:
            path = write_csv(Path(folder) / 'conocimientos_maritimo_2024_12.csv',
                             [record(JAN), record(FEB), record(AIR), record('invalid')])
            sources = inspect_sources(path)
            self.assertEqual(sources[0].rows, 4)
            groups = available_periods(sources)
            self.assertEqual(set(groups), {('2025-01', 'maritimo'), ('2025-02', 'maritimo'), ('2026-04', 'aereo')})
            self.assertEqual(groups[('2025-01', 'maritimo')][0][1].rows, 1)

    def test_nonconsecutive_periods_keep_exact_zeroes_and_audit_selection(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = write_csv(root / 'mixed.csv', [record(FEB), record(JAN), record(AIR), record(JAN, '002', courier='S')])
            before = hashlib.sha256(path.read_bytes()).hexdigest()
            selected = {('2025-01', 'maritimo'), ('2026-04', 'aereo')}
            output = root / 'base.xlsx'
            report = build_base(inspect_sources(path), output, periods=selected, audit_dir=root / 'audit')
            self.assertEqual(report['exported'], 2)
            with closing(load_workbook(output, read_only=True)) as wb:
                rows = list(wb['BLS'].values)
                self.assertEqual([row[4] for row in rows[1:]], ['2025-01', '2026-04'])
                for cells in wb['BLS'].iter_rows(min_row=2):
                    self.assertEqual(cells[1].value, '001')
                    self.assertEqual(cells[2].value, '0000001234567890123456789')
                    self.assertTrue(all(cell.data_type == 's' and cell.number_format == '@' for cell in cells))
            with closing(sqlite3.connect(report['audit'])) as db:
                filters = json.loads(db.execute("SELECT value FROM metadata WHERE key='filters'").fetchone()[0])
                self.assertEqual(filters['periods'], [list(key) for key in sorted(selected)])
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)

    def test_partial_and_empty_periods_are_visible_without_exporting_checkpoints(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = root / 'conocimientos.sqlite3'
            store = KnowledgeStore(root / 'unused.csv', {}, db_path=path)
            store.prepare_selection([(1, '', {}, JAN), (2, '', {}, FEB), (3, '', {}, AIR)])
            store.start_run()
            store.claim(JAN)
            store.save(JAN, 1, [record(JAN)], {'pages': 1, 'read': 1, 'kept': 1})
            store.save_checkpoint(FEB, GridPage({}, {}, [record(FEB)], 0, False))
            store.claim(AIR)
            store.save(AIR, 3, [], {'pages': 1, 'read': 0, 'kept': 0})
            store.close()
            sources = inspect_sources(path)
            groups = {p.key: p for p in sources[0].periods}
            self.assertEqual((groups[('2025-01', 'maritimo')].saved, groups[('2025-01', 'maritimo')].pending), (1, 0))
            self.assertEqual((groups[('2025-02', 'maritimo')].rows, groups[('2025-02', 'maritimo')].pending), (0, 1))
            self.assertEqual((groups[('2026-04', 'aereo')].rows, groups[('2026-04', 'aereo')].saved), (0, 1))
            self.assertEqual(build_base(sources, root / 'base.xlsx', audit_dir=root / 'audit')['exported'], 1)

    def test_same_output_prefers_sqlite_and_other_csv_remains_selectable(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            period = root / '2025/01/maritimo'
            path = period / 'SQLite/conocimientos_maritimo_2025_01.sqlite3'
            store = KnowledgeStore(period / 'unused.csv', {}, db_path=path)
            store.prepare_selection([(1, '', {}, JAN)])
            store.start_run()
            store.claim(JAN)
            store.save(JAN, 1, [record(JAN)], {'pages': 1, 'read': 1, 'kept': 1})
            store.close()
            companion = write_csv(period / 'CSV/conocimientos_maritimo_2025_01.csv', [record(JAN)])
            other = write_csv(root / 'other.csv', [record(FEB)])
            sources, errors = discover_sources([root])
            self.assertFalse(errors)
            choices = default_selection(sources)
            self.assertEqual({s.path for s in sources if choices[s.key]}, {path.resolve(), other.resolve()})
            self.assertFalse(choices[next(s.key for s in sources if s.path == companion.resolve())])

    def test_cancelled_discovery_does_not_report_cancellation_as_unreadable_file(self):
        with tempfile.TemporaryDirectory() as folder:
            path = write_csv(Path(folder) / 'source.csv', [record(JAN)])
            stop = Event()
            with self.assertRaises(BuildCancelled):
                discover_sources([path], stop=stop, progress=lambda event: stop.set())

    def test_empty_period_selection_preserves_existing_excel(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = write_csv(root / 'source.csv', [record(JAN)])
            output = root / 'base.xlsx'
            sources = inspect_sources(path)
            build_base(sources, output, audit_dir=root / 'audit')
            before = output.read_bytes()
            with self.assertRaises(ValueError):
                build_base(sources, output, periods=set(), audit_dir=root / 'audit')
            self.assertEqual(output.read_bytes(), before)

    def test_cli_repeatable_periods_generate_only_requested_months(self):
        from bls_base_cli import main
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            write_csv(root / 'mixed.csv', [record(JAN), record(FEB), record(AIR)])
            output = root / 'base.xlsx'
            def build(*args, **kwargs):
                return build_base(*args, audit_dir=root / 'audit', **kwargs)
            arguments = ['bls_base_cli.py', '--carpeta', str(root), '--periodo', '2025-01',
                         '--periodo', '2026-04', '--salida', str(output)]
            with patch('sys.argv', arguments), patch('bls_base_cli.build_base', side_effect=build), redirect_stdout(io.StringIO()):
                main()
            with closing(load_workbook(output, read_only=True)) as wb:
                self.assertEqual([r[4] for r in list(wb['BLS'].values)[1:]], ['2025-01', '2026-04'])


class PeriodPanelTests(unittest.TestCase):
    def test_choose_months_across_years_and_preserve_hidden_selections(self):
        if os.environ.get('BLS_PERIOD_PANEL_ISOLATED') != '1':
            env = dict(os.environ, BLS_PERIOD_PANEL_ISOLATED='1')
            env['PYTHONPATH'] = os.pathsep.join([str(Path(__file__).resolve().parent),
                                               str(Path(__file__).resolve().parents[1] / 'src')])
            result = subprocess.run([sys.executable, '-m', 'unittest',
                'test_bls_periods.PeriodPanelTests.test_choose_months_across_years_and_preserve_hidden_selections'],
                env=env, capture_output=True, text=True, timeout=20)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            return
        import ttkbootstrap as ttk
        from bls_base_panel import BLSBaseWindow
        with tempfile.TemporaryDirectory() as folder:
            root_path = Path(folder)
            write_csv(root_path / 'mixed.csv', [record(JAN), record(FEB), record(AIR)])
            root = ttk.Window(themename='darkly')
            root.withdraw()
            try:
                with patch('bls_base_panel.default_roots', return_value=[root_path]):
                    panel = BLSBaseWindow(root)
                    panel.window.withdraw()
                deadline = time.monotonic() + 5
                while panel.active and time.monotonic() < deadline:
                    root.update()
                    time.sleep(.02)
                self.assertFalse(panel.active)
                self.assertFalse(panel.sources)
                panel.scan()
                deadline = time.monotonic() + 5
                while panel.active and time.monotonic() < deadline:
                    root.update()
                    time.sleep(.02)
                self.assertFalse(panel.active)
                self.assertEqual(set(panel.filters[0]['values']), {'Todos', '2025', '2026'})
                self.assertEqual(len(panel.chosen_periods()), 3)
                items = panel.period_table.get_children()
                panel.period_table.selection_set(items[:2])
                panel._toggle_period_items(items[:2])
                self.assertEqual(set(panel.period_table.selection()), set(items[:2]))
                panel.select_periods(False)
                panel.period_choices[('2025-01', 'maritimo')] = True
                panel.period_choices[('2026-04', 'aereo')] = True
                panel.render_periods()
                panel.year.set('2025')
                panel.filter_changed()
                self.assertEqual(panel.chosen_periods(), {('2025-01', 'maritimo')})
                self.assertEqual(set(panel.filters[1]['values']), {'Todos', 'Enero', 'Febrero'})
                panel.year.set('Todos')
                panel.filter_changed()
                self.assertEqual(len(panel.chosen_periods()), 2)
                panel.output.set(str(root_path / 'base.xlsx'))
                def build(*args, **kwargs):
                    return build_base(*args, audit_dir=root_path / 'audit', **kwargs)
                with patch('bls_base_panel.build_base', side_effect=build), patch('bls_base_panel.messagebox.showinfo'):
                    panel.generate()
                    deadline = time.monotonic() + 5
                    while panel.active and time.monotonic() < deadline:
                        root.update()
                        time.sleep(.02)
                self.assertFalse(panel.active)
                with closing(load_workbook(root_path / 'base.xlsx', read_only=True)) as wb:
                    self.assertEqual([r[4] for r in list(wb['BLS'].values)[1:]], ['2025-01', '2026-04'])
                panel.close()
            finally:
                root.destroy()


if __name__ == '__main__':
    unittest.main()
