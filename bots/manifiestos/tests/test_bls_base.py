import csv
import hashlib
import os
import subprocess
import sys
import sqlite3
import tempfile
import time
import unittest
from pathlib import Path
from contextlib import closing
from threading import Event
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from bls_base import (BuildCancelled, OUTPUT_COLUMNS, Source, available_periods, build_base,
                      default_selection, discover_sources)
from conocimientos import COLUMNS, GridPage, KnowledgeStore
from conocimientos_queue import OutputLock


JAN = '1-0-20240739-P027-02/01/25'
FEB = '1-0-20250024-P027-02/02/25'
AIR = '4-0-26000124-P010-01/04/26'


def row(mid=JAN, seq='001', bl='0000123', courier='N'):
    return [mid, seq, bl, courier] + [''] * (len(COLUMNS) - 4)


class BLSBaseTests(unittest.TestCase):
    def test_sqlite_matches_xlsx_and_stores_only_text_bls_columns(self):
        # Mismo contrato: Courier, filtros, duplicados, conflictos y orden.
        a = self.csv('a.csv',[row(FEB,'2','00000999'),row(),row(JAN,'002'),row(courier='S'),row(AIR)])
        b = self.csv('b.csv',[row(),row(JAN,'004','first'),row(JAN,'004','different')])
        self.build([a,b],year=2025,mode='maritimo',sheet_rows=2)
        expected = [record for sheet in self.read() for record in sheet[1:]]
        destination = self.root/'out/base.sqlite3'
        result = build_base([a,b],destination,year=2025,mode='maritimo',sheet_rows=2,audit_dir=self.audit)
        with closing(sqlite3.connect(destination)) as db:
            self.assertEqual(db.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall(),[('BLS',)])
            columns=db.execute('PRAGMA table_info(BLS)').fetchall()
            self.assertEqual([c[1] for c in columns],OUTPUT_COLUMNS)
            self.assertTrue(all(c[2]=='TEXT' for c in columns))
            self.assertEqual(db.execute('SELECT * FROM BLS ORDER BY rowid').fetchall(),expected)
            self.assertTrue(all(isinstance(value,str) for record in expected for value in record))
        self.assertEqual(result['format'],'SQLite')
        self.assertEqual(result['tables'],1)
        self.assertEqual(result['sheets'],0)
        self.assertEqual(result['exported'],3)
        self.assertEqual(result['conflicts'],1)
        self.assertEqual(expected[0][1:3],('001','0000123'))

    def test_sqlite_cancellation_preserves_old_base_and_cleans_temporary_database(self):
        destination=self.root/'out/base.sqlite3'
        good=self.csv('good.csv',[row()])
        build_base([good],destination,audit_dir=self.audit)
        before=destination.read_bytes()
        many=self.csv('many.csv',[row(seq=str(i)) for i in range(1,2004)])
        stop=Event()
        def cancel(event):
            if event['phase']=='export': stop.set()
        with self.assertRaises(BuildCancelled):
            build_base([many],destination,audit_dir=self.audit,stop=stop,progress=cancel)
        self.assertEqual(destination.read_bytes(),before)
        self.assertFalse(list(destination.parent.glob('base.*.sqlite3')))
        self.assertFalse(list(destination.parent.glob('*-journal')))

    def test_sqlite_publication_failure_preserves_existing_base_and_releases_lock(self):
        destination=self.root/'out/base.sqlite3'
        source=self.csv('good.csv',[row()])
        build_base([source],destination,audit_dir=self.audit)
        before=destination.read_bytes()
        with patch('bls_base.os.replace',side_effect=PermissionError('Base abierta')):
            with self.assertRaises(PermissionError):
                build_base([source],destination,audit_dir=self.audit)
        self.assertEqual(destination.read_bytes(),before)
        self.assertFalse(list(destination.parent.glob('base.*.sqlite3')))
        self.assertEqual(build_base([source],destination,audit_dir=self.audit)['exported'],1)

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.output = self.root / 'out/Bls_sin_courier.xlsx'
        self.audit = self.root / 'audit'

    def tearDown(self):
        self.temp.cleanup()

    def csv(self, filename, values, delimiter=','):
        path = self.root / filename
        with path.open('w', newline='', encoding='utf-8-sig') as stream:
            writer = csv.writer(stream, delimiter=delimiter)
            writer.writerow(COLUMNS)
            writer.writerows(values)
        return Source(path, 'CSV')

    def build(self, sources, **kwargs):
        return build_base(sources, self.output, audit_dir=self.audit, **kwargs)

    def read(self):
        wb = load_workbook(self.output, read_only=True, data_only=False)
        try:
            return [list(sheet.values) for sheet in wb.worksheets]
        finally:
            wb.close()

    def test_filter_zeroes_chronology_and_unique_knowledge(self):
        a = self.csv('a.csv', [row(FEB, '2', '00999'), row(courier='S'), row(), row(AIR)])
        b = self.csv('b.csv', [row(), row(JAN, '002', '0000123')])
        report = self.build([a, b, a], year=2025, mode='maritimo')
        data = self.read()[0]
        self.assertEqual(data[0], tuple(OUTPUT_COLUMNS))
        self.assertEqual([r[4] for r in data[1:]], ['2025-01', '2025-01', '2025-02'])
        self.assertEqual(data[1], (JAN, '001', '0000123', '02/01/2025', '2025-01'))
        self.assertEqual(report['exported'], 3)
        self.assertEqual(report['duplicates'], 1)
        self.assertEqual(report['courier_excluded'], 1)
        self.assertEqual(report['filter_excluded'], 1)
        self.assertEqual(report['sources'], 2)

    def test_conflicting_bl_is_excluded_not_arbitrarily_selected(self):
        source = self.csv('conflict.csv', [row(), row(seq='1', bl='123'), row(seq='2', bl='0055')])
        report = self.build([source])
        self.assertEqual(report['conflicts'], 1)
        self.assertEqual([r[2] for r in self.read()[0][1:]], ['0055'])
        with closing(sqlite3.connect(report['audit'])) as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM provenance').fetchone()[0], 3)
            self.assertIn('contradictorio', db.execute('SELECT reason FROM issues').fetchone()[0])

    def test_xlsx_declared_padding_text_precision_and_formula_safety(self):
        path = self.root / 'source.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.title = 'Conocimientos'
        ws.append(COLUMNS)
        ws.append(row(seq=1, bl=123))
        ws['B2'].number_format = '000'
        ws['C2'].number_format = '000000'
        ws.append(row(seq=2, bl=12345678901234567))
        ws.append(row(seq=3, bl='=1+1'))
        ws.append(row(seq=4, bl=456))
        ws.append(row(seq=5, bl='00001234567890123456789'))
        wb.save(path)
        wb.close()
        report = self.build([Source(path, 'XLSX', 'Conocimientos')])
        rows = self.read()[0][1:]
        self.assertEqual(rows[0][1:3], ('001', '000123'))
        self.assertEqual(report['invalid'], 2)  # >15 dígitos y fórmula de entrada.
        self.assertEqual(report['numeric_bl'], 1)
        self.assertIn('00001234567890123456789', [r[2] for r in rows])
        # Un BL literal que empieza con '=' en CSV no se convierte en fórmula.
        literal = self.csv('literal.csv', [row(bl='=literal')])
        self.build([literal])
        wb = load_workbook(self.output)
        try:
            self.assertEqual(wb.active['C2'].value, '=literal')
            self.assertEqual(wb.active['C2'].data_type, 's')
            self.assertEqual(wb.active['C2'].number_format, '@')
        finally:
            wb.close()

    def test_sqlite_includes_only_complete_jobs_and_never_partial_pages(self):
        path = self.root / 'conocimientos.sqlite3'
        store = KnowledgeStore(self.root / 'dummy.csv', {}, db_path=path)
        store.prepare_selection([(1, '', {}, JAN), (2, '', {}, FEB)])
        store.start_run()
        store.claim(JAN)
        store.save(JAN, 1, [row(seq='1')], {'pages': 1, 'read': 1, 'kept': 1})
        store.save_checkpoint(FEB, GridPage({}, {}, [row(FEB)], 0, False))
        store.close()
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        found, errors = discover_sources([self.root])
        self.assertFalse(errors)
        sources = [s for s in found if s.kind == 'SQLite']
        self.assertEqual(len(sources), 1)
        self.assertEqual(sources[0].rows, 1)
        report = self.build(sources)
        self.assertEqual(report['exported'], 1)
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)

    def test_sqlite_filters_use_row_identity_even_with_misleading_scope(self):
        path = self.root / 'mixed.sqlite3'
        store = KnowledgeStore(self.root / 'dummy.csv', {'year': 2025, 'month': 1, 'mode': 'maritimo'}, db_path=path)
        for ordinal, mid in enumerate((JAN, FEB, AIR), 1):
            store.save(mid, ordinal, [row(mid, seq='1')], {'pages': 1, 'read': 1, 'kept': 1})
        store.close()
        report = self.build([Source(path, 'SQLite')], year=2026, month=4, mode='aereo')
        self.assertEqual(report['exported'], 1)
        self.assertEqual(self.read()[0][1][0], AIR)

    def test_discovery_overlapping_folders_headers_empty_and_other_flows(self):
        good = self.csv('conocimientos_maritimo_2025_01.csv', [row()], ';')
        self.csv('empty.csv', [])
        self.csv('fallidos.csv', [])
        (self.root / 'other.csv').write_text('BL original,IDManifiesto\n123,456\n', encoding='utf-8')
        with closing(sqlite3.connect(self.root / 'other.sqlite3')) as db:
            db.execute('CREATE TABLE unrelated (value TEXT)')
            db.commit()
        found, errors = discover_sources([self.root, self.root, good.path])
        self.assertEqual([s.path for s in found], [good.path])
        self.assertEqual(found[0].period, '2025-01')
        self.assertEqual(found[0].mode, 'maritimo')
        self.assertFalse(errors)
        self.build(found)
        self.assertEqual(self.read()[0][1][2], '0000123')

    def test_output_can_be_read_by_bls_bot_contract(self):
        self.build([self.csv('input.csv', [row()])])
        import pandas as pd
        data = pd.read_excel(self.output, sheet_name='BLS', dtype=str)
        self.assertEqual(data['BL original'][0], '0000123')
        self.assertEqual(data['Secuencia Tica'][0], '001')
        self.assertEqual(data['Mes'][0], '2025-01')
        self.assertEqual(pd.to_datetime(data['Fecha Manifiesto'], dayfirst=True)[0].year, 2025)

    def test_month_and_mode_are_filtered_per_row_not_filename(self):
        source = self.csv('conocimientos_maritimo_2025_01.csv', [row(), row(FEB), row(AIR)])
        report = self.build([source], year=2026, month=4, mode='aereo')
        self.assertEqual(report['exported'], 1)
        self.assertEqual(self.read()[0][1][0], AIR)

    def test_excel_split_keeps_header_and_all_rows(self):
        source = self.csv('many.csv', [row(seq=str(i), bl=f'000{i}') for i in range(1, 6)])
        report = self.build([source], sheet_rows=2)
        sheets = self.read()
        self.assertEqual(report['sheets'], 3)
        self.assertEqual([len(s) for s in sheets], [3, 3, 2])
        self.assertEqual([r[2] for s in sheets for r in s[1:]], [f'000{i}' for i in range(1, 6)])

    def test_empty_filter_and_bad_source_preserve_existing_output(self):
        good = self.csv('good.csv', [row()])
        self.build([good])
        original = self.output.read_bytes()
        with self.assertRaises(ValueError) as error:
            self.build([good], year=2026)
        self.assertEqual(self.output.read_bytes(), original)
        self.assertTrue(Path(error.exception.audit).exists())
        bad = self.csv('bad.csv', [row()])
        bad.path.write_text('wrong,headers\n1,2\n', encoding='utf-8')
        with self.assertRaises(ValueError):
            self.build([good, bad])
        self.assertEqual(self.output.read_bytes(), original)

    def test_cancellation_during_export_leaves_no_temporary_output(self):
        source = self.csv('many.csv', [row(seq=str(i)) for i in range(1, 2003)])
        stop = Event()
        def cancel(event):
            if event['phase'] == 'export':
                stop.set()
        with self.assertRaises(BuildCancelled) as error:
            self.build([source], stop=stop, progress=cancel)
        self.assertFalse(self.output.exists())
        self.assertFalse(list(self.output.parent.glob('*.xlsx')))
        with closing(sqlite3.connect(error.exception.audit)) as db:
            self.assertEqual(db.execute("SELECT value FROM metadata WHERE key='status'").fetchone()[0], 'cancelled')

    def test_blocked_export_preserves_previous_xlsx_and_releases_lock(self):
        source = self.csv('input.csv', [row()])
        self.build([source])
        before = self.output.read_bytes()
        with patch('bls_base.os.replace', side_effect=PermissionError('Archivo abierto en Excel')):
            with self.assertRaises(PermissionError):
                self.build([source])
        self.assertEqual(before, self.output.read_bytes())
        self.assertFalse(list(self.output.parent.glob('Bls_sin_courier.*.xlsx')))
        self.assertEqual(self.build([source])['exported'], 1)

    def test_invalid_ids_and_records_are_audited(self):
        source = self.csv('invalid.csv', [row(), row('wrong'), row(seq='0'), row(bl=''), row('1-0-20250001-P001-31/02/25')])
        report = self.build([source])
        self.assertEqual(report['invalid'], 4)
        self.assertEqual(report['exported'], 1)

    def test_source_is_never_overwritten(self):
        source = self.csv('input.csv', [row()])
        with self.assertRaises(ValueError):
            build_base([source], source.path, audit_dir=self.audit)
        with source.path.open(encoding='utf-8-sig') as stream:
            self.assertEqual(len(list(csv.reader(stream))), 2)

    def test_same_output_cannot_be_generated_by_two_instances(self):
        source = self.csv('input.csv', [row()])
        self.output.parent.mkdir(parents=True)
        lock = OutputLock(str(self.output) + '.lock')
        try:
            with self.assertRaises(RuntimeError):
                self.build([source])
        finally:
            lock.close()
        self.assertEqual(self.build([source])['exported'], 1)

    def test_real_periods_are_detected_from_mixed_rows(self):
        self.csv('conocimientos_maritimo_2025_01.csv', [row(), row(FEB), row(AIR)])
        sources, errors = discover_sources([self.root])
        self.assertFalse(errors)
        self.assertEqual(set(available_periods(sources)),
                         {('2025-01', 'maritimo'), ('2025-02', 'maritimo'), ('2026-04', 'aereo')})

    def test_multiple_period_selection_across_years(self):
        source = self.csv('mixed.csv', [row(), row(FEB), row(AIR)])
        report = self.build([source], periods={('2025-01', 'maritimo'), ('2026-04', 'aereo')})
        self.assertEqual(report['exported'], 2)
        self.assertEqual([r[4] for r in self.read()[0][1:]], ['2025-01', '2026-04'])

    def test_prefer_sqlite_to_its_export_but_keep_other_csv(self):
        source = self.csv('conocimientos.csv', [row(seq='1')])
        other = self.csv('other.csv', [row(FEB)])
        store = KnowledgeStore(source.path, {}, db_path=source.path.with_suffix('.sqlite3'), accept_legacy=True)
        store.save(JAN, 1, [row(seq='1')], {'pages': 1, 'read': 1, 'kept': 1})
        store.close()
        sources, errors = discover_sources([self.root])
        self.assertFalse(errors)
        choices = default_selection(sources)
        self.assertFalse(choices[source.key])
        self.assertTrue(choices[other.key])
        self.assertTrue(choices[(str(source.path.with_suffix('.sqlite3').resolve()).casefold(), '')])

    def test_pending_only_period_is_visible_but_not_selected_as_data(self):
        store = KnowledgeStore(self.root / 'dummy.csv', {}, db_path=self.root / 'pending.sqlite3')
        store.prepare_selection([(1, '', {}, JAN)])
        store.close()
        sources, errors = discover_sources([self.root])
        self.assertFalse(errors)
        self.assertEqual(len(sources), 1)
        period = sources[0].periods[0]
        self.assertEqual((period.saved, period.pending, period.rows), (0, 1, 0))
        self.assertFalse(default_selection(sources)[sources[0].key])


class BLSPanelTests(unittest.TestCase):
    def test_compact_window_keeps_generate_and_output_visible_at_small_sizes(self):
        if self._run_isolated():
            return
        import ttkbootstrap as ttk
        from bls_base_panel import BLSBaseWindow
        ttk.Style.instance = None
        root = ttk.Window(themename='darkly')
        root.withdraw()
        panel = None
        try:
            with patch('bls_base_panel.default_roots', return_value=[]), \
                 patch.object(ttk.Toplevel, 'winfo_screenwidth', return_value=860), \
                 patch.object(ttk.Toplevel, 'winfo_screenheight', return_value=700):
                panel = BLSBaseWindow(root)
            deadline = time.monotonic() + 5
            while panel.active and time.monotonic() < deadline:
                root.update()
                time.sleep(.01)
            self.assertFalse(panel.active)
            root.update()
            self.assertLessEqual(panel.window.winfo_width(), 800)
            self.assertLessEqual(panel.window.winfo_height(), 600)
            self.assertEqual(panel.tabs.select(), str(panel.months_tab))
            def assert_visible(widget):
                x = widget.winfo_rootx() - panel.window.winfo_rootx()
                y = widget.winfo_rooty() - panel.window.winfo_rooty()
                self.assertTrue(widget.winfo_viewable())
                self.assertGreaterEqual(x, 0)
                self.assertGreaterEqual(y, 0)
                self.assertLessEqual(x + widget.winfo_width(), panel.window.winfo_width())
                self.assertLessEqual(y + widget.winfo_height(), panel.window.winfo_height())
            for scale in (1.33, 2.0):
                root.tk.call('tk', 'scaling', scale)
                for size in ('800x600', '680x500', '860x650'):
                    panel.window.geometry(size)
                    for tab in (panel.months_tab, panel.files_tab):
                        panel.tabs.select(tab)
                        root.update()
                        for widget in (panel.build_button, panel.stop_button, panel.output_combo,
                                       panel.output_entry, panel.output_button, panel.scan_button,
                                       panel.folder_button, panel.diagnostic_button, *panel.filters):
                            assert_visible(widget)
                        self.assertGreater(panel.output_entry.winfo_width(), 80)
            self.assertEqual(panel.build_button.cget('text'), 'Generar base')
        finally:
            if panel is not None:
                panel.close()
            root.destroy()

    def test_one_app_origin_external_inputs_and_sqlite_output_selection(self):
        if self._run_isolated(): return
        import ttkbootstrap as ttk
        from bls_base_panel import BLSBaseWindow,APP_INPUT,EXTERNAL_INPUT
        ttk.Style.instance=None
        with tempfile.TemporaryDirectory() as folder:
            base=Path(folder)
            app=base/'app'
            external=base/'external'
            for directory,mid in ((app,JAN),(external,FEB)):
                directory.mkdir()
                with (directory/'source.csv').open('w',newline='',encoding='utf-8') as stream:
                    writer=csv.writer(stream)
                    writer.writerow(COLUMNS)
                    writer.writerow(row(mid))
            root=ttk.Window(themename='darkly')
            root.withdraw()
            def wait(panel):
                deadline=time.monotonic()+5
                while panel.active and time.monotonic()<deadline:
                    root.update()
                    time.sleep(.02)
                self.assertFalse(panel.active)
            try:
                with patch('bls_base_panel.default_roots',return_value=[app]):
                    panel=BLSBaseWindow(root)
                    panel.window.withdraw()
                    wait(panel)
                self.assertEqual(panel.input_source.get(),APP_INPUT)
                self.assertNotIn(str(app),panel.roots_text.get())
                panel.input_source.set(EXTERNAL_INPUT)
                panel.source_changed()
                wait(panel)
                self.assertEqual(panel.roots,[])
                with patch('bls_base_panel.filedialog.askdirectory',return_value=str(external)):
                    panel.add_folder()
                    wait(panel)
                self.assertEqual(panel.roots,[external.resolve()])
                panel.input_source.set(APP_INPUT)
                panel.source_changed()
                wait(panel)
                self.assertEqual(panel.roots,[app])
                panel.scan()
                wait(panel)
                panel.output.set(str(base/'out/base.xlsx'))
                panel.output_format.set('SQLite (.sqlite3)')
                panel.output_format_changed()
                self.assertEqual(Path(panel.output.get()).suffix,'.sqlite3')
                def build(*args,**kwargs):
                    return build_base(*args,audit_dir=base/'audit',**kwargs)
                with patch('bls_base_panel.build_base',side_effect=build),patch('bls_base_panel.messagebox.showinfo'):
                    panel.generate()
                    wait(panel)
                with closing(sqlite3.connect(panel.output.get())) as db:
                    self.assertEqual(db.execute('SELECT * FROM BLS').fetchone(),
                                     (JAN,'001','0000123','02/01/2025','2025-01'))
                self.assertIn('SQLite',panel.status.get())
                panel.close()
            finally:
                root.destroy()

    def test_choose_scope_before_reading_and_reload_when_month_changes(self):
        if self._run_isolated():
            return
        import ttkbootstrap as ttk
        from bls_base_panel import BLSBaseWindow
        from resultados import output_paths
        ttk.Style.instance = None
        with tempfile.TemporaryDirectory() as folder:
            paths = []
            for month, mid in ((1,JAN),(2,FEB)):
                location = output_paths({'output_manifest_conocimientos_folder':folder},
                                        'conocimientos','maritimo',month,2025)
                store = KnowledgeStore(location.csv, {}, db_path=location.database)
                store.save(mid,1,[row(mid)],{'kept':1})
                store.close()
                location.csv.write_bytes(b'\xff CSV companion should not be opened')
                paths.append(location)
            root = ttk.Window(themename='darkly')
            root.withdraw()
            def wait(panel):
                deadline=time.monotonic()+5
                while panel.active and time.monotonic()<deadline:
                    root.update()
                    time.sleep(.02)
                self.assertFalse(panel.active)
            try:
                with patch('bls_base_panel.default_roots', return_value=[Path(folder)]), \
                        patch('bls_base_panel.discover_sources', wraps=discover_sources) as read:
                    panel = BLSBaseWindow(root)
                    panel.window.withdraw()
                    wait(panel)
                    read.assert_not_called()
                    self.assertFalse(panel.sources)
                    panel.year.set('2025')
                    panel.mode.set('Marítimo')
                    panel.month.set('Enero')
                    panel.filter_changed()
                    panel.scan()
                    wait(panel)
                    self.assertEqual(read.call_args.args[0],[paths[0].database.resolve()])
                    self.assertTrue(panel.ready_to_build())
                    self.assertEqual(panel.chosen_periods(),{('2025-01','maritimo')})
                    panel.month.set('Febrero')
                    panel.filter_changed()
                    self.assertFalse(panel.ready_to_build())
                    with patch('bls_base_panel.messagebox.showwarning') as warning:
                        panel.generate()
                        warning.assert_called_once()
                    panel.scan()
                    wait(panel)
                    self.assertEqual(panel.chosen_periods(),{('2025-02','maritimo')})
                    loaded = list(panel.sources.values())
                    panel.month.set('Todos')
                    panel.filter_changed()
                    with patch('bls_base_panel.discover_sources', return_value=(loaded,[(str(paths[0].database),'database is locked')])):
                        panel.scan()
                        wait(panel)
                    self.assertFalse(panel.ready_to_build())
                    panel.month.set('Febrero')
                    panel.filter_changed()
                    self.assertTrue(panel.ready_to_build())
                    panel.close()
            finally:
                root.destroy()

    def _run_isolated(self):
        # Evitar que el GC de los trabajadores asyncio destruya Tk de estas pruebas.
        if os.environ.get('BLS_BASE_PANEL_ISOLATED') == '1':
            return False
        env = dict(os.environ, BLS_BASE_PANEL_ISOLATED='1')
        env['PYTHONPATH'] = os.pathsep.join([str(Path(__file__).resolve().parent),
                                           str(Path(__file__).resolve().parents[1]/'src')])
        result = subprocess.run([sys.executable, '-m', 'unittest', self.id()],
                                env=env, capture_output=True, text=True, timeout=20)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        return True

    def test_gui_selects_real_periods_and_ignores_hidden_years(self):
        if self._run_isolated():
            return
        import ttkbootstrap as ttk
        from bls_base_panel import BLSBaseWindow
        ttk.Style.instance = None
        with tempfile.TemporaryDirectory() as folder:
            with (Path(folder) / 'mixed.csv').open('w', encoding='utf-8', newline='') as stream:
                writer = csv.writer(stream)
                writer.writerow(COLUMNS)
                writer.writerows([row(), row(FEB), row(AIR)])
            root = ttk.Window(themename='darkly')
            root.withdraw()
            try:
                with patch('bls_base_panel.default_roots', return_value=[Path(folder)]):
                    panel = BLSBaseWindow(root)
                    panel.window.withdraw()
                deadline = time.monotonic() + 5
                while panel.active and time.monotonic() < deadline:
                    root.update()
                    time.sleep(.02)
                self.assertFalse(panel.active)
                self.assertFalse(panel.sources)
                panel.scan()
                deadline = time.monotonic() + 5
                while panel.active and time.monotonic() < deadline:
                    root.update()
                    time.sleep(.02)
                self.assertFalse(panel.active)
                self.assertEqual(len(panel.chosen_periods()), 3)
                self.assertEqual(tuple(panel.filters[0]['values']), ('Todos', '2025', '2026'))
                panel.year.set('2025')
                panel.mode.set('Marítimo')
                panel.filter_changed()
                self.assertEqual(panel.chosen_periods(), {('2025-01', 'maritimo'), ('2025-02', 'maritimo')})
                panel.select_periods(False)
                self.assertFalse(panel.chosen_periods())
                panel.period_choices[('2025-01', 'maritimo')] = True
                panel.render_periods()
                self.assertEqual(panel.chosen_periods(), {('2025-01', 'maritimo')})
                panel.year.set('2026')
                panel.mode.set('Aéreo')
                panel.filter_changed()
                self.assertEqual(panel.chosen_periods(), {('2026-04', 'aereo')})
                panel.close()
            finally:
                root.destroy()

    def test_main_app_has_a_separate_reusable_bls_section(self):
        if self._run_isolated():
            return
        import ttkbootstrap as ttk
        import manifest_gui
        ttk.Style.instance = None
        root = ttk.Window(themename='darkly')
        root.withdraw()
        def walk(widget):
            for child in widget.winfo_children():
                yield child
                yield from walk(child)
        def mainloop():
            buttons = [w for w in walk(root) if isinstance(w, ttk.Button) and w.cget('text') == 'Crear base para BLs_BOT']
            self.assertEqual(len(buttons), 1)
            buttons[0].invoke()
            buttons[0].invoke()
        try:
            with patch('manifest_gui.organize_results', return_value=[]), patch('manifest_gui.ttk.Window', return_value=root), patch.object(root, 'mainloop', side_effect=mainloop), patch('manifest_gui.BLSBaseWindow') as panel:
                manifest_gui.main()
                panel.assert_called_once_with(root)
                panel.return_value.window.lift.assert_called_once()
        finally:
            root.destroy()

    def test_background_scan_selection_filters_and_diagnostics(self):
        if self._run_isolated():
            return
        import ttkbootstrap as ttk
        from bls_base_panel import BLSBaseWindow
        ttk.Style.instance = None
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'conocimientos.csv'
            with path.open('w', encoding='utf-8', newline='') as stream:
                writer = csv.writer(stream)
                writer.writerow(COLUMNS)
                writer.writerow(row())
            root = ttk.Window(themename='darkly')
            root.withdraw()
            try:
                with patch('bls_base_panel.default_roots', return_value=[Path(folder)]):
                    panel = BLSBaseWindow(root)
                    panel.window.withdraw()
                deadline = time.monotonic() + 5
                while panel.active and time.monotonic() < deadline:
                    root.update()
                    time.sleep(.02)
                self.assertFalse(panel.active)
                self.assertFalse(panel.sources)
                panel.scan()
                deadline = time.monotonic() + 5
                while panel.active and time.monotonic() < deadline:
                    root.update()
                    time.sleep(.02)
                self.assertFalse(panel.active)
                self.assertEqual(len(panel.sources), 1)
                panel.select_all(False)
                self.assertFalse(any(panel.selected.values()))
                panel.select_all(True)
                self.assertTrue(all(panel.selected.values()))
                panel.year.set('2025')
                panel.mode.set('Marítimo')
                panel.suggest_output()
                self.assertTrue(panel.output.get().endswith('Bls_sin_courier_maritimo_2025.xlsx'))
                panel.output.set(str(Path(folder) / 'base.xlsx'))
                def build(*args, **kwargs):
                    return build_base(*args, audit_dir=Path(folder) / 'audit', **kwargs)
                with patch('bls_base_panel.build_base', side_effect=build), patch('bls_base_panel.messagebox.showinfo') as info:
                    panel.generate()
                    deadline = time.monotonic() + 5
                    while panel.active and time.monotonic() < deadline:
                        root.update()
                        time.sleep(.02)
                    self.assertFalse(panel.active)
                    self.assertTrue(Path(panel.output.get()).exists())
                    self.assertIsNotNone(panel.last_audit)
                    info.assert_called_once()
                panel.diagnostics()
                panel.close()
            finally:
                root.destroy()


if __name__ == '__main__':
    unittest.main()
