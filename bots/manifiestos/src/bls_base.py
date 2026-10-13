"""Construcción local de bases BL: fuentes de solo lectura y XLSX de texto."""
import csv
import json
import math
import os
import re
import sqlite3
import tempfile
import uuid
from contextlib import closing
from dataclasses import dataclass
from datetime import datetime
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell

from conocimientos import COLUMNS
from conocimientos_queue import OutputLock
from resultados import PROJECT_ROOT, output_base, application_base, csv_for_database, parse_result_name

OUTPUT_COLUMNS = ['IDManifiesto', 'Secuencia Tica', 'BL original', 'Fecha Manifiesto', 'Mes']
REQUIRED = OUTPUT_COLUMNS[:3] + ['Courier']
MODES = {'1': 'maritimo', '4': 'aereo', '7': 'terrestre'}
MAX_SHEET_ROWS = 1_048_575  # Más una fila de encabezado.
OUTPUT_FORMATS = {'.xlsx': 'Excel', '.sqlite3': 'SQLite', '.sqlite': 'SQLite', '.db': 'SQLite'}
EXPORT_QUERY = '''SELECT manifest,sequence,bl,fecha,period FROM entries
    WHERE conflict=0 ORDER BY period,date_iso,mode,manifest_number,manifest,sequence_key'''


class BuildCancelled(Exception):
    pass


class NotKnowledgeSource(ValueError):
    pass


@dataclass(frozen=True)
class DownloadedPeriod:
    period: str
    mode: str
    rows: int
    saved: int | None = None
    pending: int | None = None

    @property
    def key(self):
        return self.period, self.mode


@dataclass(frozen=True)
class Source:
    path: Path
    kind: str
    sheet: str = ''
    period: str = ''
    mode: str = ''
    rows: int | None = None
    periods: tuple[DownloadedPeriod, ...] = ()

    @property
    def key(self):
        return str(self.path.resolve()).casefold(), self.sheet


@dataclass(frozen=True)
class SourceCandidate:
    path: Path
    kind: str
    period: str = ''
    mode: str = ''
    historical: bool = False
    managed: bool = False


def catalog_sources(roots, stop=None, include_history=False):
    """Enumera nombres/rutas; no abre bases ni recorre filas de CSV/Excel."""
    paths, errors = set(), []
    excluded = {'fallidos', 'control', 'archivos_originales', '__pycache__'}
    if not include_history:
        excluded.add('historicos')
    for root in roots:
        check_stop(stop)
        root = Path(root).resolve()
        if not root.exists():
            errors.append((str(root), 'La carpeta no existe'))
            continue
        if root.is_file():
            paths.add(root)
            continue
        for directory, folders, files in os.walk(root):
            check_stop(stop)
            folders[:] = [name for name in folders if name.casefold() not in excluded]
            paths.update(Path(directory)/name for name in files)
    found = []
    for path in sorted(paths):
        check_stop(stop)
        kind = {'.csv':'CSV', '.sqlite3':'SQLite', '.sqlite':'SQLite', '.db':'SQLite', '.xlsx':'XLSX'}.get(path.suffix.lower())
        if not kind or path.name.startswith('~$'):
            continue
        parsed = parse_result_name(path)
        if parsed and (parsed[0] != 'conocimientos' or parsed[-1]):
            continue
        historical = 'historicos' in {p.casefold() for p in path.parts} or '_historico_' in path.stem.casefold()
        if historical and not include_history:
            continue
        period, mode = _hint(path)
        # Solo las carpetas del bot garantizan el período del nombre. Un archivo
        # externo puede mezclar meses: se verifica su contenido al cargarlo.
        parts = tuple(p.casefold() for p in path.parts)
        managed = bool(parsed and any(parts[i:i+3] == ('conocimientos', str(parsed[3]), parsed[1])
                                     for i in range(len(parts)-2)))
        found.append(SourceCandidate(path, kind, period, mode, historical, managed))
    return found, errors


def choose_candidates(candidates, year=None, month=None, mode=None, kind='preferred'):
    chosen = [c for c in candidates if (not c.managed or
        ((year is None or c.period[:4] == str(year)) and
         (month is None or c.period[5:] == f'{month:02d}') and
         (mode is None or c.mode == mode))) and
        (kind in ('preferred','all') or c.kind == kind)]
    if kind == 'preferred':
        companions = set()
        for candidate in chosen:
            if candidate.kind == 'SQLite':
                try:
                    companions.add(csv_for_database(candidate.path).resolve())
                except (OSError, ValueError, KeyError):
                    pass  # La lectura del archivo dará el diagnóstico correspondiente.
        chosen = [c for c in chosen if c.kind != 'CSV' or c.path.resolve() not in companions]
    return chosen


def check_stop(stop):
    if stop is not None and stop.is_set():
        raise BuildCancelled('Generación detenida. El Excel anterior permanece intacto.')


def headers(values):
    return [str(v or '').replace('\ufeff', '').replace('\xa0', ' ').strip() for v in values]


def cell_text(cell):
    """Solo reconstruye padding si el XLSX declara explícitamente ese ancho."""
    value = cell.value
    if value is None:
        return ''
    if cell.data_type in ('f', 'e'):
        raise ValueError('Identificador con fórmula o error: requiere un valor literal')
    if isinstance(value, bool):
        raise ValueError('Identificador booleano')
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or value != int(value) or value < 0:
            raise ValueError('Identificador numérico no entero')
        text = str(int(value))
        if len(text) > 15:
            raise ValueError('Identificador numérico de más de 15 cifras: posible pérdida de precisión en Excel')
        fmt = cell.number_format or ''
        return text.zfill(len(fmt)) if re.fullmatch(r'0+', fmt) else text
    return str(value).strip()


def default_roots(config):
    public = output_base(config, 'conocimientos')
    paths = [public / 'conocimientos',
             application_base(config, 'conocimientos') / 'SQLite' / 'conocimientos']
    return list(dict.fromkeys(path.resolve() for path in paths if path.is_dir()))


def _csv_reader(stream):
    sample = stream.read(8192)
    stream.seek(0)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
    except csv.Error:
        dialect = csv.excel
    return csv.reader(stream, dialect)


def readonly(path):
    return sqlite3.connect(Path(path).resolve().as_uri() + '?mode=ro', uri=True, timeout=5)


def _sqlite_query(db):
    names = {row[0] for row in db.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    if not {'rows', 'manifests', 'metadata'} <= names:
        raise NotKnowledgeSource('No es una base de Conocimientos')
    scope = json.loads(db.execute('SELECT scope FROM metadata').fetchone()[0])
    if scope.get('flow', 'conocimientos') != 'conocimientos':
        raise NotKnowledgeSource('La base pertenece a otro procedimiento')
    query = ' FROM rows r JOIN manifests m ON m.id=r.manifest'
    if 'jobs' in names:
        # Las bases históricas pueden tener resultados confirmados antes de la cola.
        query += " LEFT JOIN jobs j ON j.id=r.manifest WHERE (j.state='Guardado' OR j.id IS NULL)"
    return scope, query


def _hint(path):
    parsed = parse_result_name(path)
    if parsed:
        _, mode, month, year, _ = parsed
        return f'{year:04d}-{month:02d}', mode
    match = re.search(r'(maritimo|aereo|terrestre)_(\d{4})_(\d{2})', path.stem, re.I)
    return (f'{match[2]}-{match[3]}', match[1].lower()) if match else ('', '')


def manifest_identity(mid):
    parts = str(mid).strip().split('-')
    if len(parts) != 5 or parts[0] not in MODES or parts[1] not in ('0', '1'):
        raise ValueError('IDManifiesto incompatible')
    if not re.fullmatch(r'\d+', parts[2]) or not re.fullmatch(r'[A-Za-z]\d+', parts[3]):
        raise ValueError('Número o ubicación del manifiesto inválidos')
    return parts, datetime.strptime(parts[4], '%d/%m/%y').date()


def _period_key(mid, cache):
    mid = str(mid or '').strip()
    if mid not in cache:
        try:
            parts, day = manifest_identity(mid)
            cache[mid] = day.strftime('%Y-%m'), MODES[parts[0]]
        except ValueError:
            cache[mid] = None
    return cache[mid]


def _file_periods(records, stop=None, progress=None):
    counts, cache, total = {}, {}, 0
    for index, mid in enumerate(records):
        if index % 2000 == 0:
            check_stop(stop)
            if progress:
                progress(index)
        total += 1
        key = _period_key(mid, cache)
        if key is not None:
            counts[key] = counts.get(key, 0) + 1
    return total, tuple(DownloadedPeriod(*key, count) for key, count in sorted(counts.items()))


def _sqlite_periods(db, query, stop=None):
    counts, states, cache = {}, {}, {}
    total = 0
    for mid, count in db.execute('SELECT r.manifest,COUNT(*)' + query + ' GROUP BY r.manifest'):
        check_stop(stop)
        total += count
        key = _period_key(mid, cache)
        if key is not None:
            counts[key] = counts.get(key, 0) + count
    names = {r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    if 'jobs' in names:
        for mid, state in db.execute('SELECT id,state FROM jobs'):
            check_stop(stop)
            key = _period_key(mid, cache)
            if key is not None:
                saved, pending = states.get(key, (0, 0))
                states[key] = saved + (state == 'Guardado'), pending + (state != 'Guardado')
    return tuple(DownloadedPeriod(*key, counts.get(key, 0), *states.get(key, (None, None)))
                 for key in sorted(counts.keys() | states.keys())), total


def inspect_sources(path, stop=None, progress=None):
    check_stop(stop)
    path = Path(path).resolve()
    period, mode = _hint(path)
    if path.suffix.lower() in ('.sqlite3','.sqlite','.db'):
        with closing(readonly(path)) as db:
            last_update = time.monotonic()
            def interrupted():
                nonlocal last_update
                if stop is not None and stop.is_set():
                    return 1
                if progress and time.monotonic() - last_update > .25:
                    progress(None)
                    last_update = time.monotonic()
                return 0
            db.set_progress_handler(interrupted, 2000)
            db.execute('BEGIN')
            scope, query = _sqlite_query(db)
            first = db.execute('SELECT r.data' + query + ' LIMIT 1').fetchone()
            if first and len(json.loads(first[0])) != len(COLUMNS):
                return []
            try:
                periods, count = _sqlite_periods(db, query, stop)
            except sqlite3.OperationalError:
                check_stop(stop)
                raise
            if not periods:
                return []
            if scope.get('year') and scope.get('month'):
                period = f"{int(scope['year']):04d}-{int(scope['month']):02d}"
            return [Source(path, 'SQLite', period=period, mode=scope.get('mode', mode), rows=count, periods=periods)]
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as stream:
            reader = _csv_reader(stream)
            names = headers(next(reader, []))
            if set(REQUIRED) <= set(names) and len(names) == len(set(names)):
                mid_index = names.index('IDManifiesto')
                count, periods = _file_periods((row[mid_index] if len(row) > mid_index else ''
                    for row in reader if any(v.strip() for v in row)), stop, progress)
                if count:
                    return [Source(path, 'CSV', period=period, mode=mode, rows=count, periods=periods)]
        return []
    if path.suffix.lower() == '.xlsx':
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            found = []
            for sheet in wb.worksheets:
                iterator = sheet.iter_rows()
                names = headers([c.value for c in next(iterator, [])])
                if set(REQUIRED) <= set(names) and len(names) == len(set(names)):
                    mid_index = names.index('IDManifiesto')
                    count, periods = _file_periods((row[mid_index].value if len(row) > mid_index else ''
                        for row in iterator if any(c.value is not None for c in row)), stop, progress)
                    if count:
                        found.append(Source(path, 'XLSX', sheet.title, period, mode, count, periods))
            return found
        finally:
            wb.close()
    return []


def discover_sources(roots, stop=None, progress=None, *, include_history=True, prefer_sqlite=False):
    """Detecta por esquema, no por una ruta personal o un nombre fijo."""
    candidates, errors = catalog_sources(roots, stop, include_history)
    candidates = choose_candidates(candidates, kind='preferred' if prefer_sqlite else 'all')
    paths, sources = [c.path for c in candidates], []
    for index, path in enumerate(paths, 1):
        check_stop(stop)
        if progress:
            progress({'phase': 'scan', 'current': index, 'total': len(paths), 'file': str(path)})
        try:
            def report_rows(count):
                if progress:
                    progress({'phase':'scan', 'current':index, 'total':len(paths), 'file':str(path), 'rows':count})
            sources.extend(inspect_sources(path, stop, report_rows))
        except BuildCancelled:
            raise
        except NotKnowledgeSource:
            continue
        except Exception as exc:
            check_stop(stop)
            errors.append((str(path), str(exc)))
    check_stop(stop)
    sources.sort(key=lambda s: (s.period, s.mode, s.path.name, s.sheet, s.kind))
    return sources, errors


def available_periods(sources):
    """Períodos reales agrupados, conservando sus fuentes y estados."""
    groups = {}
    for source in sources:
        for period in source.periods:
            groups.setdefault(period.key, []).append((source, period))
    return dict(sorted(groups.items()))


def default_selection(sources):
    """Prefiere SQLite al CSV exportado de la misma salida."""
    databases = {s.path.resolve() for s in sources if s.kind == 'SQLite' and s.rows}
    exported = {csv_for_database(path).resolve() for path in databases}
    selected = {}
    for source in sources:
        counterpart = (source.path.parent.parent / 'SQLite' / (source.path.stem + '.sqlite3')
                       if source.path.parent.name == 'CSV' else source.path.with_suffix('.sqlite3'))
        selected[source.key] = bool(source.rows) and not (source.kind == 'CSV' and
            (counterpart.resolve() in databases or source.path.resolve() in exported))
    return selected


def iter_records(source, year=None, month=None, mode=None):
    if source.kind == 'SQLite':
        db = readonly(source.path)
        try:
            db.execute('BEGIN')
            _, query = _sqlite_query(db)
            conditions, parameters = [], []
            if mode:
                conditions.append('r.manifest LIKE ?')
                parameters.append(next(k for k, v in MODES.items() if v == mode) + '-%')
            if year is not None:
                conditions.append('substr(r.manifest,-2)=?')
                parameters.append(f'{year % 100:02d}')
            if month is not None:
                conditions.append('substr(r.manifest,-5,2)=?')
                parameters.append(f'{month:02d}')
            if conditions:
                query += (' AND ' if ' WHERE ' in query else ' WHERE ') + ' AND '.join(conditions)
            for index, (data, mid) in enumerate(db.execute('SELECT r.data,r.manifest' + query + ' ORDER BY m.ordinal,r.sequence', parameters), 1):
                row = json.loads(data)
                if len(row) != len(COLUMNS) or row[0] != mid:
                    yield index, None, 'Columnas o identidad SQLite incompatibles'
                else:
                    yield index, dict(zip(COLUMNS, row)), ''
        finally:
            db.close()
    elif source.kind == 'CSV':
        with source.path.open(encoding='utf-8-sig', newline='') as stream:
            reader = _csv_reader(stream)
            names = headers(next(reader, []))
            if not set(REQUIRED) <= set(names) or len(names) != len(set(names)):
                raise ValueError('Encabezados CSV incompatibles o duplicados')
            for index, row in enumerate(reader, 2):
                if not any(v.strip() for v in row):
                    continue
                if len(row) != len(names):
                    yield index, None, 'Fila CSV incompleta o con columnas adicionales'
                else:
                    yield index, dict(zip(names, row)), ''
    else:
        wb = load_workbook(source.path, read_only=True, data_only=False)
        try:
            iterator = wb[source.sheet].iter_rows()
            names = headers([c.value for c in next(iterator, [])])
            if not set(REQUIRED) <= set(names) or len(names) != len(set(names)):
                raise ValueError('Encabezados XLSX incompatibles o duplicados')
            for index, row in enumerate(iterator, 2):
                if not any(c.value is not None for c in row):
                    continue
                try:
                    record = {name: cell_text(row[i]) for i, name in enumerate(names) if name in REQUIRED}
                    bl_cell = row[names.index('BL original')]
                    record['_numeric_bl'] = isinstance(bl_cell.value, (int, float)) and not re.fullmatch(r'0+', bl_cell.number_format or '')
                    yield index, record, ''
                except (ValueError, IndexError) as exc:
                    yield index, None, str(exc)
        finally:
            wb.close()


def normalize_record(record):
    mid, sequence, bl = (str(record.get(name, '')).strip() for name in OUTPUT_COLUMNS[:3])
    parts, day = manifest_identity(mid)
    if not re.fullmatch(r'\d+', sequence) or not 1 <= int(sequence) <= 2**63 - 1:
        raise ValueError('Secuencia Tica no válida')
    if not bl or bl.lower() in ('none', 'nan', 'null'):
        raise ValueError('BL original vacío')
    # El número puede pertenecer al año previo; el período es la fecha de arribo.
    return (mid, int(sequence), sequence, bl, day.strftime('%d/%m/%Y'),
            day.strftime('%Y-%m'), day.isoformat(), MODES[parts[0]], int(parts[2]))


def export_sqlite(db, destination, stats, count, stop, progress):
    """Publica solo BLS: todas las columnas TEXT; auditoría separada y orden estable."""
    fd, temporary = tempfile.mkstemp(prefix=destination.stem+'.', suffix='.sqlite3', dir=destination.parent)
    os.close(fd)
    try:
        with closing(sqlite3.connect(temporary)) as output:
            columns = ','.join('"'+name.replace('"','""')+'" TEXT NOT NULL' for name in OUTPUT_COLUMNS)
            output.execute('CREATE TABLE BLS ('+columns+')')
            batch = []
            for row in db.execute(EXPORT_QUERY):
                check_stop(stop)
                batch.append(tuple(str(value) for value in row))
                stats['exported'] += 1
                if len(batch) == 2000:
                    output.executemany('INSERT INTO BLS VALUES (?,?,?,?,?)',batch)
                    batch.clear()
                    if progress:
                        progress({'phase':'export','current':stats['exported'],'total':count,**stats})
            output.executemany('INSERT INTO BLS VALUES (?,?,?,?,?)',batch)
            check_stop(stop)
            output.commit()
        stats['tables'] = 1
        return temporary
    except BaseException:
        if os.path.exists(temporary):
            os.unlink(temporary)
        raise


def export_xlsx(db, destination, stats, count, stop, progress, sheet_rows):
    wb = Workbook(write_only=True)
    sheet, temporary = None, None
    try:
        for row_index, row in enumerate(db.execute(EXPORT_QUERY)):
            check_stop(stop)
            if row_index % sheet_rows == 0:
                stats['sheets'] += 1
                sheet = wb.create_sheet('BLS' if stats['sheets'] == 1 else f"BLS_{stats['sheets']:03d}")
                sheet.freeze_panes = 'A2'
                for col, width in zip('ABCDE',(38,19,32,22,14)):
                    sheet.column_dimensions[col].width = width
                sheet.append(OUTPUT_COLUMNS)
                sheet.auto_filter.ref = f'A1:E{min(sheet_rows,count-row_index)+1}'
            cells = []
            for value in row:
                cell = WriteOnlyCell(sheet,value=str(value))
                cell.data_type, cell.number_format = 's','@'
                cells.append(cell)
            sheet.append(cells)
            stats['exported'] += 1
            if progress and stats['exported'] % 2000 == 0:
                progress({'phase':'export','current':stats['exported'],'total':count,**stats})
        check_stop(stop)
        fd, temporary = tempfile.mkstemp(prefix=destination.stem+'.',suffix='.xlsx',dir=destination.parent)
        os.close(fd)
        wb.save(temporary)
        return temporary
    except BaseException:
        if temporary and os.path.exists(temporary):
            os.unlink(temporary)
        raise
    finally:
        for worksheet in wb.worksheets:
            if not worksheet.closed:
                worksheet.close()
            if worksheet._writer and os.path.exists(worksheet._writer.out):
                worksheet._writer.cleanup()
        wb.close()


def build_base(sources, destination, *, year=None, mode=None, month=None,
               periods=None, audit_dir=None, stop=None, progress=None, sheet_rows=MAX_SHEET_ROWS):
    """Unifica por manifiesto+secuencia; discrepancias de BL se excluyen y auditan."""
    destination = Path(destination).resolve()
    output_format = OUTPUT_FORMATS.get(destination.suffix.lower())
    if not output_format:
        raise ValueError('La salida debe ser .xlsx o SQLite (.sqlite3, .sqlite, .db)')
    if mode not in (None, *MODES.values()) or (month is not None and not 1 <= month <= 12):
        raise ValueError('Filtro de modalidad o mes inválido')
    if not sources or not 1 <= sheet_rows <= MAX_SHEET_ROWS:
        raise ValueError('Seleccione al menos una fuente y un tamaño de hoja válido')
    selected_periods = None if periods is None else frozenset(tuple(key) for key in periods)
    if selected_periods is not None:
        if not selected_periods or any(len(key) != 2 or key[1] not in MODES.values() or
                not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', key[0]) for key in selected_periods):
            raise ValueError('Seleccione al menos un período y modalidad válidos')
    unique = {source.key: source for source in sources}
    if any(s.path.resolve() == destination for s in unique.values()):
        raise ValueError('La salida no puede sustituir un archivo fuente')
    destination.parent.mkdir(parents=True, exist_ok=True)
    audit_dir = Path(audit_dir or application_base({}) / 'bases_bls/auditoria').resolve()
    audit_dir.mkdir(parents=True, exist_ok=True)
    audit = audit_dir / f"base_{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:8]}.sqlite3"
    stats = {'read': 0, 'courier_excluded': 0, 'filter_excluded': 0, 'duplicates': 0,
             'invalid': 0, 'conflicts': 0, 'numeric_bl': 0, 'exported': 0, 'sources': len(unique), 'sheets': 0,
             'output': str(destination), 'audit': str(audit), 'format': output_format, 'tables': 0}
    lock = OutputLock(str(destination) + '.lock')
    db, temp = None, None
    try:
        db = sqlite3.connect(audit)
        db.executescript('''
            CREATE TABLE metadata (key TEXT PRIMARY KEY,value TEXT NOT NULL);
            CREATE TABLE sources (id INTEGER PRIMARY KEY,path TEXT,sheet TEXT,kind TEXT);
            CREATE TABLE entries (
                manifest TEXT, sequence_key INTEGER, sequence TEXT, bl TEXT,
                fecha TEXT, period TEXT, date_iso TEXT, mode TEXT, manifest_number INTEGER,
                conflict INTEGER NOT NULL DEFAULT 0, PRIMARY KEY(manifest,sequence_key));
            CREATE TABLE provenance (source_id INTEGER,row_number INTEGER,manifest TEXT,sequence_key INTEGER,bl TEXT);
            CREATE TABLE issues (source_id INTEGER,row_number INTEGER,reason TEXT);
        ''')
        db.execute('INSERT INTO metadata VALUES (?,?)', ('filters', json.dumps({'year': year, 'mode': mode,
            'month': month, 'periods': None if selected_periods is None else sorted(selected_periods)})))
        for source_id, source in enumerate(unique.values(), 1):
            check_stop(stop)
            db.execute('INSERT INTO sources VALUES (?,?,?,?)', (source_id, str(source.path), source.sheet, source.kind))
            for row_number, record, error in iter_records(source, year, month, mode):
                check_stop(stop)
                stats['read'] += 1
                if not error and str(record.get('Courier', '')).strip().upper() != 'N':
                    stats['courier_excluded'] += 1
                    continue
                try:
                    if error:
                        raise ValueError(error)
                    values = normalize_record(record)
                except (ValueError, TypeError) as exc:
                    stats['invalid'] += 1
                    db.execute('INSERT INTO issues VALUES (?,?,?)', (source_id, row_number, str(exc)))
                    continue
                mid, seq, _, bl, _, period, date_iso, row_mode, _ = values
                if ((year is not None and int(period[:4]) != year) or
                    (month is not None and int(period[5:]) != month) or
                    (mode is not None and row_mode != mode) or
                    (selected_periods is not None and (period, row_mode) not in selected_periods)):
                    stats['filter_excluded'] += 1
                    continue
                if record.get('_numeric_bl'):
                    stats['numeric_bl'] += 1
                    db.execute('INSERT INTO issues VALUES (?,?,?)', (source_id, row_number,
                        'BL numérico sin formato de ceros: se conserva el valor disponible; no se pueden inferir ceros perdidos'))
                db.execute('INSERT INTO provenance VALUES (?,?,?,?,?)', (source_id, row_number, mid, seq, bl))
                previous = db.execute('SELECT bl,conflict FROM entries WHERE manifest=? AND sequence_key=?', (mid, seq)).fetchone()
                if previous is None:
                    db.execute('INSERT INTO entries VALUES (?,?,?,?,?,?,?,?,?,0)', values)
                elif previous[0] == bl:
                    stats['duplicates'] += 1
                else:
                    if not previous[1]:
                        stats['conflicts'] += 1
                    db.execute('UPDATE entries SET conflict=1 WHERE manifest=? AND sequence_key=?', (mid, seq))
                    db.execute('INSERT INTO issues VALUES (?,?,?)',
                               (source_id, row_number, f'BL contradictorio para {mid}, secuencia {seq}: {previous[0]!r} / {bl!r}'))
                if stats['read'] % 2000 == 0:
                    db.commit()
                    if progress:
                        progress({'phase': 'read', 'current': source_id, 'total': len(unique), 'file': str(source.path), **stats})
            db.commit()
            if progress:
                progress({'phase': 'read', 'current': source_id, 'total': len(unique), 'file': str(source.path), **stats})
        count = db.execute('SELECT COUNT(*) FROM entries WHERE conflict=0').fetchone()[0]
        if not count:
            raise ValueError('No hay BL válidos Courier=N para los filtros seleccionados. Revise las fuentes y el diagnóstico.')
        if output_format == 'SQLite':
            temp = export_sqlite(db,destination,stats,count,stop,progress)
        else:
            temp = export_xlsx(db,destination,stats,count,stop,progress,sheet_rows)
        check_stop(stop)
        os.replace(temp, destination)
        temp = None
        db.execute('INSERT OR REPLACE INTO metadata VALUES (?,?)', ('status', 'completed'))
        db.execute('INSERT OR REPLACE INTO metadata VALUES (?,?)', ('summary', json.dumps(stats, ensure_ascii=False)))
        db.commit()
        return stats
    except BaseException as exc:
        exc.audit = str(audit)
        if db is not None:
            db.execute('INSERT OR REPLACE INTO metadata VALUES (?,?)', ('status', 'cancelled' if isinstance(exc, BuildCancelled) else 'failed'))
            db.execute('INSERT OR REPLACE INTO metadata VALUES (?,?)', ('error', str(exc)))
            db.execute('INSERT OR REPLACE INTO metadata VALUES (?,?)', ('summary', json.dumps(stats, ensure_ascii=False)))
            db.commit()
        raise
    finally:
        if db is not None:
            db.close()
        if temp and os.path.exists(temp):
            os.unlink(temp)
        lock.close()
