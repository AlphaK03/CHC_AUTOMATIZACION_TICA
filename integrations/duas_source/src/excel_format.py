"""Workbook layout verified against the user's legacy September 2026 files."""
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
import os
import shutil
import tempfile
from zipfile import ZipFile
from output_paths import filesystem_path

BASE_COLUMNS = ('aduana', 'ano', 'numero', 'despacho', 'ingreso', 'declarante',
                'bultos', 'modalidad', 'plazo', 'adu_destino')


def empty_columns(regimen):
    return BASE_COLUMNS + (('ubi_destino',) if regimen in ('EXPORTACION', 'TRANSITO', 'ZONA FRANCA') else ())


def write_excel(path, sheet, columns, rows):
    """Legacy pandas layout without pandas; keep all values as literal text."""
    path = filesystem_path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook(write_only=True)
    ws = book.create_sheet(sheet)
    header = []
    side = Side(style='thin')
    for column in columns:
        cell = WriteOnlyCell(ws, value=column)
        cell.data_type = 's'
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='top')
        cell.border = Border(left=side, right=side, top=side, bottom=side)
        header.append(cell)
    ws.append(header)
    count = 0
    try:
        for values in rows:
            cells = []
            for value in values:
                cell = WriteOnlyCell(ws, value='' if value is None else str(value))
                cell.data_type = 's'
                cells.append(cell)
            ws.append(cells)
            count += 1
        book.save(path)
    finally:
        book.close()
    # Write-only worksheets omit their dimensions. Legacy files include them,
    # and consumers using read_only/max_row need that metadata as well.
    fd, temporary = tempfile.mkstemp(suffix='.xlsx', dir=path.parent)
    os.close(fd)
    try:
        with ZipFile(path) as source, ZipFile(temporary, 'w') as target:
            for entry in source.infolist():
                with source.open(entry) as incoming, target.open(entry, 'w') as outgoing:
                    if entry.filename == 'xl/worksheets/sheet1.xml':
                        prefix = incoming.read(4096)
                        # SpreadsheetML requires sheetPr BEFORE dimension, then
                        # sheetViews. Putting dimension first passes openpyxl's
                        # reader but causes Microsoft Excel to reject the file.
                        end = prefix.index(b'<sheetViews')
                        dimension = f'<dimension ref="A1:{get_column_letter(len(columns))}{count + 1}"/>'.encode()
                        outgoing.write(prefix[:end] + dimension + prefix[end:])
                    shutil.copyfileobj(incoming, outgoing)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    return count
