"""Fresh destinations, portable paths and Windows paths beyond MAX_PATH."""
import os
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from data_store import RangeStore, REGIMENES
from exports import publish, publish_originals, atomic_json, safe_target
from output_paths import output_path, filesystem_path, display_path, prepare_output_folder
from openpyxl import load_workbook


class OutputPathTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name).resolve()
        self.addCleanup(self.temp.cleanup)
        def clean_long_paths():
            target = filesystem_path(self.root)
            assert display_path(target) == self.root
            shutil.rmtree(target)
        self.addCleanup(clean_long_paths)

    def test_default_relative_quoted_and_environment_paths(self):
        app = self.root / 'portable app'
        self.assertEqual(output_path('', app), app / 'DUAs_2025')
        self.assertEqual(output_path('"nueva carpeta/mes"', app), app / 'nueva carpeta' / 'mes')
        self.assertEqual(output_path(str(self.root / 'another folder'), app), self.root / 'another folder')
        with patch.dict(os.environ, {'DUA_EXPORT_TEST': str(self.root)}):
            self.assertEqual(output_path('$DUA_EXPORT_TEST/resultados', app), self.root / 'resultados')

    def test_fresh_nested_destination_is_created_for_both_export_types(self):
        base = self.root / 'nueva carpeta' / 'año 2026' / 'septiembre'
        self.assertFalse(base.exists())
        paths = publish(base, '2026-09', [(1, 15)], [[], [], []])
        self.assertTrue(all(Path(p).exists() for p in paths))
        store = RangeStore(self.root / 'test.sqlite3')
        try:
            key = store.plan('2026-09', ['TRANSITO'], [(1, 15)])[0]
            store.complete(key, [])
            paths, _ = publish_originals(base / 'otra ubicacion', '2026-09', store, 'halves')
            self.assertEqual(len(paths), 1)
            self.assertTrue(Path(paths[0]).exists())
        finally:
            store.close()

    @unittest.skipUnless(os.name == 'nt', 'Windows extended paths')
    def test_long_paths_support_publication_backup_mode_changes_and_recovery(self):
        base = self.root / ('a' * 90) / ('b' * 90)
        store = RangeStore(self.root / 'test.sqlite3')
        try:
            keys = store.plan('2026-09', REGIMENES, [(1, 15), (16, 30)])
            for key in keys:
                store.complete(key, [])
            lists, _ = store.collect('2026-09', [(1, 15), (16, 30)])
            totals = publish(base, '2026-09', [(1, 15), (16, 30)], lists, store=store)
            halves, _ = publish_originals(base, '2026-09', store, 'halves')
            self.assertEqual(len(halves), 18)
            self.assertTrue(any(len(path) > 260 for path in halves))
            for path in [*totals, *halves]:
                self.assertFalse(path.startswith('\\\\?\\'))
                wb = load_workbook(filesystem_path(path), read_only=True)
                try:
                    self.assertEqual(wb.sheetnames, ['septiembre26'])
                finally:
                    wb.close()
            before = filesystem_path(totals[0]).read_bytes()
            import exports
            copy = exports.atomic_copy
            def fail(source, target):
                if Path(target).name.endswith('Importacion.xlsx'):
                    raise PermissionError('locked workbook')
                copy(source, target)
            with patch('exports.atomic_copy', side_effect=fail):
                with self.assertRaises(PermissionError):
                    publish(base, '2026-09', [(1, 15)], [[], [], []], store=store)
            self.assertEqual(filesystem_path(totals[0]).read_bytes(), before)
            self.assertFalse(filesystem_path(base / exports.JOURNAL).exists())
        finally:
            store.close()

    def test_destination_file_and_denied_access_have_clear_errors(self):
        path = self.root / 'existing_file'
        path.write_text('unchanged')
        with self.assertRaisesRegex(OSError, 'un archivo, no a una carpeta'):
            prepare_output_folder(path)
        with patch('output_paths.tempfile.TemporaryFile', side_effect=PermissionError()):
            with self.assertRaisesRegex(OSError, 'No hay permiso'):
                prepare_output_folder(self.root / 'no permission')
        self.assertEqual(path.read_text(), 'unchanged')

    def test_metadata_creates_parents_and_path_guards_still_apply(self):
        target = self.root / 'new' / 'nested' / 'state.json'
        atomic_json(target, {'ok': True})
        self.assertTrue(target.exists())
        with self.assertRaises(ValueError):
            safe_target(self.root, '../outside.xlsx')
        self.assertEqual(display_path(filesystem_path(target)), target)
