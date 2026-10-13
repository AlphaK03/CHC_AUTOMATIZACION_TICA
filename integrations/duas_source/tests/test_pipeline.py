"""Pruebas locales: clasificación original, cobertura, reanudación y publicación."""
from datetime import date
import json
from pathlib import Path
import sqlite3
import subprocess
import sys
import tempfile
from threading import Event
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from openpyxl import load_workbook
from data_store import RangeStore, RunLock, REGIMENES, REGIMENES_IMPORTACION, OUTPUT_NAMES, period_for, ranges_for
import exports
from main import unir_archivos_finales
from runner import run_downloads
from tica_dua_bot import parse_xml, parse_date, TICADUAExtractor, Cancelled
from captcha import lookup, word_for


def record(number, year='2026'):
    return {'aduana': '001', 'ano': year, 'numero': f'{number:06d}', 'modalidad': 'prueba'}


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.path = self.root / 'duas.sqlite3'
        self.store = RangeStore(self.path)
        self.addCleanup(self.store.close)

    def fill(self, ranges=((1, 15), (16, 31)), period='2026-03', empty=False):
        ids = self.store.plan(period, REGIMENES, ranges)
        for key in ids:
            query = self.store.get(key)
            number = REGIMENES.index(query['regimen']) + 1
            self.store.complete(key, [] if empty else [record(number), record(number)])
        return ids

    def read_output(self, path):
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            self.assertEqual(wb.sheetnames, ['marzo26'])
            rows = list(wb.active.iter_rows(values_only=True))
            self.assertEqual(rows[0], ('NroDUA',))
            return [row[0] for row in rows[1:]]
        finally:
            wb.close()

    def test_original_classification_and_three_excel_contracts(self):
        self.assertEqual(REGIMENES, ['EXPORTACION', 'EXPORTACION TEMPORAL', 'IMPORTACION DEFINITIVA',
            'IMPORTACION TEMPORAL', 'PERFECCIONAMIENTO ACTIVO', 'REIMPORTACION',
            'REIMPORTACION PERFECCIONAMIENTO PASIVO', 'TRANSITO', 'ZONA FRANCA'])
        self.assertEqual(REGIMENES_IMPORTACION, {'IMPORTACION DEFINITIVA', 'IMPORTACION TEMPORAL',
            'REIMPORTACION', 'REIMPORTACION PERFECCIONAMIENTO PASIVO'})
        self.fill()
        files, missing = unir_archivos_finales('Marzo', '2026', 'Ambas partes', lambda _: None, self.root, self.path)
        self.assertEqual(missing, [])
        self.assertEqual([Path(f).name for f in files], list(OUTPUT_NAMES))
        self.assertEqual(self.read_output(files[0]), [f'001-2026-{i:06d}' for i in range(1, 10)])
        self.assertEqual(self.read_output(files[1]), ['001-2026-000008'])
        self.assertEqual(self.read_output(files[2]), [f'001-2026-{i:06d}' for i in (3, 4, 6, 7)])

    def test_partial_month_missing_regime_never_changes_previous_exports(self):
        for name in OUTPUT_NAMES:
            (self.root / name).write_bytes(b'previous data')
        self.fill(ranges=((1, 15),))
        files, missing = unir_archivos_finales('Marzo', '2026', 'Ambas partes', lambda _: None, self.root, self.path)
        self.assertFalse(files)
        self.assertEqual(len(missing), 9)
        self.assertTrue(all((self.root / name).read_bytes() == b'previous data' for name in OUTPUT_NAMES))

    def test_year_and_month_are_isolated_legacy_files_not_trusted(self):
        self.fill(period='2025-03')
        for regimen in REGIMENES:
            folder = self.root / regimen
            folder.mkdir()
            (folder / f'DUAS_MARZO_{regimen}_1-15.xlsx').write_bytes(b'old ambiguous file')
        self.assertEqual(len(self.store.collect('2026-03', [(1, 15)])[1]), 9)
        self.assertEqual(len(self.store.collect('2025-04', [(1, 15)])[1]), 9)

    def test_custom_ranges_prove_coverage_without_repeating_duas(self):
        self.fill(ranges=((1, 7), (8, 15)))
        lists, missing = self.store.collect('2026-03', [(1, 15)])
        self.assertEqual(missing, [])
        self.assertEqual(len(lists[0]), 9)

    def test_cannot_export_half_of_a_larger_range_without_dates(self):
        self.fill(ranges=((1, 31),))
        self.assertEqual(len(self.store.collect('2026-03', [(1, 15)])[1]), 9)
        self.assertFalse(self.store.collect('2026-03', [(1, 15), (16, 31)])[1])

    def test_refresh_failure_not_hidden_by_overlapping_success(self):
        self.fill(ranges=((1, 15), (1, 7), (8, 15)))
        query = self.store.plan('2026-03', ['TRANSITO'], [(1, 7)], refresh=True)[0]
        self.store.fail(query, 'portal failed')
        self.assertEqual(len(self.store.collect('2026-03', [(1, 15)])[1]), 1)

    def test_empty_verified_ranges_export_three_empty_files(self):
        self.fill(empty=True)
        paths, missing = unir_archivos_finales('Marzo', '2026', 'Ambas partes', lambda _: None, self.root, self.path)
        self.assertFalse(missing)
        self.assertEqual(len(paths), 3)
        self.assertTrue(all(self.read_output(path) == [] for path in paths))

    def test_transaction_rollback_preserves_previous_records(self):
        key = self.fill()[0]
        before = self.store.get(key)
        with self.assertRaises(ValueError):
            self.store.complete(key, [record(99), {'aduana': 'bad'}])
        self.assertEqual(self.store.get(key), before)
        self.assertEqual(self.store.db.execute('SELECT dua FROM records WHERE query_id=?', (key,)).fetchone()[0], '001-2026-000001')

    def test_database_rejects_duplicate_query_and_dua(self):
        key = self.fill()[0]
        with self.assertRaises(sqlite3.IntegrityError):
            with self.store.transaction() as db:
                db.execute('INSERT INTO records VALUES (?,?,?)', (key, '001-2026-000001', '{}'))
        self.assertEqual(self.store.get(key)['count'], 1)

    def test_lock_rejects_another_writer(self):
        with RunLock(self.path):
            with self.assertRaises(RuntimeError):
                with RunLock(self.path):
                    pass

    def test_publication_failure_restores_entire_previous_set(self):
        self.fill()
        lists, _ = self.store.collect('2026-03', [(1, 15), (16, 31)])
        exports.publish(self.root, '2026-03', [(1, 31)], lists)
        names = (*OUTPUT_NAMES, exports.MANIFEST)
        previous = {name: (self.root / name).read_bytes() for name in names}
        copy = exports.atomic_copy
        def fail(source, target):
            if Path(target).name == OUTPUT_NAMES[1]:
                raise PermissionError('Excel abierto')
            copy(source, target)
        with patch('exports.atomic_copy', side_effect=fail):
            with self.assertRaises(PermissionError):
                exports.publish(self.root, '2026-04', [(1, 30)], [[], [], []])
        self.assertEqual({name: (self.root / name).read_bytes() for name in names}, previous)
        self.assertFalse((self.root / exports.JOURNAL).exists())

    def test_process_crash_recovers_previous_exports(self):
        previous = {}
        for name in (*OUTPUT_NAMES, exports.MANIFEST):
            previous[name] = b'previous-' + name.encode()
            (self.root / name).write_bytes(previous[name])
        code = '''import sys,os
sys.path.insert(0,sys.argv[1])
import exports
from pathlib import Path
copy=exports.atomic_copy
def crash(source,target):
    copy(source,target)
    os._exit(9)
exports.atomic_copy=crash
exports.publish(Path(sys.argv[2]),'2026-03',[(1,15)],[[],[],[]])
'''
        result = subprocess.run([sys.executable, '-c', code, str(Path(__file__).resolve().parents[1] / 'src'), str(self.root)], timeout=20)
        self.assertEqual(result.returncode, 9)
        self.assertTrue((self.root / exports.JOURNAL).exists())
        exports.recover_publication(self.root)
        self.assertEqual({name: (self.root / name).read_bytes() for name in previous}, previous)

    def test_date_ranges_and_captcha_match_existing_catalog(self):
        self.assertEqual(period_for('Marzo', '2026'), '2026-03')
        self.assertEqual(ranges_for('2024-02', 'Ambos'), [(1, 15), (16, 29)])
        with self.assertRaises(ValueError):
            ranges_for('2026-02', 'Rango personalizado', 16, 30)
        self.assertEqual(parse_date('16/03/2026'), date(2026, 3, 16))
        self.assertEqual(parse_date('2026-03-16'), date(2026, 3, 16))
        wb = load_workbook(Path(__file__).resolve().parents[1] / 'Names.xlsx', read_only=True, data_only=True)
        try:
            rows = wb.active.iter_rows(values_only=True)
            header = next(rows)
            expected = {str(int(r[header.index('Number')])): str(r[header.index('Word')]).strip() for r in rows}
        finally:
            wb.close()
        self.assertEqual(dict(lookup()), expected)
        self.assertEqual(word_for('https://example.test/Captcha/images/74.jpg'), expected['74'])

    def test_xml_requires_complete_valid_document(self):
        self.assertEqual(parse_xml('<body><root /></body>'), [])
        self.assertEqual(parse_xml('<root><dua><aduana>1</aduana><ano>2026</ano><numero>8</numero></dua></root>')[0]['numero'], '000008')
        for raw in ['<root><dua>', '<body>Sin respuesta</body>', '<root><error>fallo</error></root>', '<root>Server error</root>', '<root><dua/></root>']:
            with self.assertRaises(ValueError):
                parse_xml(raw)

    def test_runner_retries_only_failed_ranges_and_resumes(self):
        calls, closes = [], []
        class Browser:
            def __init__(self, *args, **kwargs): pass
            def fetch_range(self, regimen, period, start, end):
                calls.append(start)
                if start == 1 and calls.count(1) == 1:
                    raise RuntimeError('portal error')
                return [record(start)]
            def close(self): closes.append(1)
        cancel = Event()
        with patch.object(cancel, 'wait', return_value=False):
            pending = run_downloads(self.path, {}, ['TRANSITO'], 'Marzo', '2026', 'Ambos', lambda _: None, cancel, factory=Browser)
        self.assertEqual(pending, 0)
        self.assertEqual(calls, [1, 16, 1, 8])
        run_downloads(self.path, {}, ['TRANSITO'], 'Marzo', '2026', 'Ambos', lambda _: None, cancel, factory=Browser)
        self.assertEqual(calls, [1, 16, 1, 8])
        self.assertEqual(len(closes), 2)

    def test_cancel_and_storage_failure_preserve_retryable_ranges(self):
        cancel = Event()
        class Browser:
            def __init__(self, *args, **kwargs): pass
            def fetch_range(self, *args): raise Cancelled()
            def close(self): pass
        run_downloads(self.path, {}, ['TRANSITO'], 'Marzo', '2026', 'Ambos', lambda _: None, cancel, factory=Browser)
        rows = self.store.db.execute('SELECT status FROM queries').fetchall()
        self.assertEqual([r[0] for r in rows], ['pending', 'pending'])
        cancel.clear()
        with patch.object(Browser, 'fetch_range', return_value=[record(1)]), patch.object(RangeStore, 'complete', side_effect=OSError('disk full')):
            with self.assertRaises(OSError):
                run_downloads(self.path, {}, ['TRANSITO'], 'Marzo', '2026', 'Ambos', lambda _: None, cancel, factory=Browser)
        self.store.plan('2026-03', ['TRANSITO'], [(1, 15), (16, 31)])
        self.assertEqual(self.store.db.execute("SELECT COUNT(*) FROM queries WHERE status='succeeded'").fetchone()[0], 0)


if __name__ == '__main__':
    unittest.main()
