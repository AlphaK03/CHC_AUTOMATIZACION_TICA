"""Original Excel contracts, two export modes, isolation and recovery."""
from pathlib import Path
import calendar
import json
import os
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from openpyxl import load_workbook
from data_store import RangeStore, REGIMENES, OUTPUT_NAMES
from excel_format import empty_columns
from exports import publish_originals, publish, MANIFEST
import exports
from main import exportar_originales


def record(regimen, number):
    row = dict.fromkeys(empty_columns(regimen), '')
    row.update(aduana='001', ano='2026', numero=f'{number:06d}', despacho='01/09/2026',
               ingreso='01/09/2026', declarante='000012345', bultos='001', modalidad='003')
    return row


class OriginalExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.database = self.root / 'test.sqlite3'
        self.store = RangeStore(self.database)
        self.addCleanup(self.store.close)

    def fill(self, period='2026-09', regimens=REGIMENES, ranges=((1, 7), (8, 15), (16, 22), (23, 30))):
        ids = self.store.plan(period, regimens, ranges)
        for key in ids:
            query = self.store.get(key)
            self.store.complete(key, [record(query['regimen'], query['day_start']), record(query['regimen'], 99)])
        return ids

    def test_two_halves_hide_retry_fragments_and_match_legacy_layout(self):
        self.fill()
        paths, pending = publish_originals(self.root, '2026-09', self.store, 'halves')
        self.assertFalse(pending)
        self.assertEqual(len(paths), 18)
        for path in paths:
            path = Path(path)
            regimen = path.parent.parent.name
            self.assertEqual(path.parent.name, '2026-09')
            self.assertTrue(path.name.endswith(('_1-15.xlsx', '_16-30.xlsx')))
            wb = load_workbook(path)
            try:
                self.assertEqual(wb.sheetnames, ['septiembre26'])
                ws = wb.active
                self.assertEqual([c.value for c in ws[1]], list(empty_columns(regimen)))
                self.assertEqual(ws.max_row, 4)  # shared DUA is not duplicated
                self.assertIsNone(ws.freeze_panes)
                self.assertIsNone(ws.auto_filter.ref)
                self.assertFalse(ws.column_dimensions)
                self.assertFalse(ws.tables)
                for cell in ws[1]:
                    self.assertTrue(cell.font.bold)
                    self.assertEqual(cell.alignment.horizontal, 'center')
                    self.assertEqual(cell.alignment.vertical, 'top')
                    self.assertEqual(cell.border.bottom.style, 'thin')
                for cell in ws[2][:9]:
                    if cell.value is not None:
                        self.assertEqual(cell.data_type, 's')
                        self.assertEqual(cell.number_format, 'General')
                self.assertEqual(ws['A2'].value, '001')
            finally:
                wb.close()
            wb = load_workbook(path, read_only=True)
            try:
                self.assertEqual(wb.active.max_row, 4)
                self.assertEqual(wb.active.max_column, len(empty_columns(regimen)))
            finally:
                wb.close()

    def test_month_mode_replaces_managed_halves_but_keeps_other_years_and_user_files(self):
        self.fill()
        halves, _ = publish_originals(self.root, '2026-09', self.store, 'halves')
        manual = Path(halves[0]).parent / 'mis_notas.xlsx'
        manual.write_bytes(b'personal workbook')
        self.fill('2025-09')
        older, _ = publish_originals(self.root, '2025-09', self.store, 'month')
        older_content = {p: Path(p).read_bytes() for p in older}
        paths, _ = publish_originals(self.root, '2026-09', self.store, 'month')
        self.assertEqual(len(paths), 9)
        self.assertTrue(all(Path(p).name.endswith('_2026.xlsx') for p in paths))
        self.assertTrue(all(not Path(p).exists() for p in halves))
        self.assertEqual(manual.read_bytes(), b'personal workbook')
        self.assertEqual({p: Path(p).read_bytes() for p in older}, older_content)
        self.assertTrue(all(not (self.root / name).exists() for name in OUTPUT_NAMES))

    def test_pending_day_allows_only_completed_half_and_never_incomplete_month(self):
        self.fill(regimens=['TRANSITO'], ranges=((1, 7), (8, 15), (16, 22)))
        paths, pending = exportar_originales('Septiembre', '2026', 'halves', lambda _: None, self.root, self.database)
        self.assertEqual(len(paths), 1)
        self.assertTrue(paths[0].endswith('_1-15.xlsx'))
        self.assertEqual(pending['TRANSITO'], list(range(23, 31)))
        monthly, _ = publish_originals(self.root, '2026-09', self.store, 'month')
        self.assertFalse(monthly)

    def test_failed_refresh_removes_stale_current_export_but_history_remains(self):
        ids = self.fill(regimens=['TRANSITO'], ranges=((1, 15), (16, 30)))
        paths, _ = publish_originals(self.root, '2026-09', self.store, 'month')
        self.store.fail(ids[1], 'timeout')
        current, _ = publish_originals(self.root, '2026-09', self.store, 'month')
        self.assertFalse(current)
        self.assertFalse(Path(paths[0]).exists())
        self.assertTrue(list((self.root / 'entregas' / '2026-09').rglob(Path(paths[0]).name)))

    def test_whole_range_is_not_invented_into_half_months(self):
        self.fill(regimens=['TRANSITO'], ranges=((1, 30),))
        paths, pending = publish_originals(self.root, '2026-09', self.store, 'halves')
        self.assertFalse(paths)
        self.assertEqual(pending['TRANSITO'], list(range(1, 31)))

    def test_empty_ranges_month_lengths_and_literal_text(self):
        for period in ('2026-02', '2024-02', '2026-03'):
            year, month = map(int, period.split('-'))
            last = calendar.monthrange(year, month)[1]
            keys = self.store.plan(period, ['TRANSITO'], [(1, 15), (16, last)])
            for key in keys:
                self.store.complete(key, [])
            paths, _ = publish_originals(self.root, period, self.store, 'halves')
            self.assertEqual(len(paths), 2)
            self.assertTrue(paths[-1].endswith(f'_16-{last}.xlsx'))
        key = self.store.plan('2026-09', ['TRANSITO'], [(1, 15)])[0]
        row = record('TRANSITO', 1)
        row['declarante'] = '=1+1'
        self.store.complete(key, [row])
        paths, _ = publish_originals(self.root, '2026-09', self.store, 'halves')
        wb = load_workbook(paths[0])
        try:
            self.assertEqual(wb.active['F2'].value, '=1+1')
            self.assertEqual(wb.active['F2'].data_type, 's')
        finally:
            wb.close()

    def test_original_failure_rolls_back_all_three_totals_and_originals(self):
        self.fill()
        lists, _ = self.store.collect('2026-09', [(1, 15), (16, 30)])
        publish(self.root, '2026-09', [(1, 15), (16, 30)], lists, store=self.store)
        manifest = 'entregas/originales-2026-09.json'
        originals = json.loads((self.root / manifest).read_text(encoding='utf-8'))['files']
        names = [*OUTPUT_NAMES, MANIFEST, manifest, *originals]
        before = {name: (self.root / name).read_bytes() for name in names}
        copy = exports.atomic_copy
        def fail(source, target):
            if Path(target).name.startswith('DUAS_'):
                raise PermissionError('workbook open')
            return copy(source, target)
        with patch('exports.atomic_copy', side_effect=fail):
            with self.assertRaises(PermissionError):
                publish(self.root, '2026-09', [(1, 15)], [[], [], []], store=self.store)
        self.assertEqual({name: (self.root / name).read_bytes() for name in names}, before)
        self.assertFalse((self.root / exports.JOURNAL).exists())

    def test_failure_after_removing_prior_mode_restores_it(self):
        self.fill(regimens=['TRANSITO'])
        paths, _ = publish_originals(self.root, '2026-09', self.store, 'month')
        old = Path(paths[0]).read_bytes()
        write = exports.atomic_json
        def fail(path, value):
            if Path(path).name == exports.JOURNAL and value.get('complete'):
                raise OSError('simulated interruption after obsolete file removal')
            return write(path, value)
        with patch('exports.atomic_json', side_effect=fail):
            with self.assertRaises(OSError):
                publish_originals(self.root, '2026-09', self.store, 'halves')
        self.assertEqual(Path(paths[0]).read_bytes(), old)
        self.assertEqual(list(Path(paths[0]).parent.glob('*.xlsx')), [Path(paths[0])])

    @unittest.skipUnless(os.environ.get('DUA_LEGACY_EXCEL'), 'Optional legacy reference folder')
    def test_cells_and_styles_against_user_reference_files(self):
        reference = Path(os.environ['DUA_LEGACY_EXCEL'])
        for regimen in REGIMENES:
            path = reference / regimen / f'DUAS_SEPTIEMBRE_{regimen}_1-15.xlsx'
            old = load_workbook(path, read_only=True)
            try:
                rows = old.active.iter_rows()
                header, first = next(rows), next(rows)
                row = {c.value: value.value or '' for c, value in zip(header, first)}
                key = self.store.plan('2026-09', [regimen], [(1, 15)])[0]
                self.store.complete(key, [row])
                paths, _ = publish_originals(self.root, '2026-09', self.store, 'halves')
                generated = next(p for p in paths if Path(p).parent.parent.name == regimen)
                new = load_workbook(generated, read_only=True)
                try:
                    exported = new.active.iter_rows()
                    for previous, actual in zip((header, first), (next(exported), next(exported))):
                        self.assertEqual([c.value for c in previous], [c.value for c in actual])
                        for a, b in zip(previous, actual):
                            self.assertEqual(a.font, b.font)
                            self.assertEqual(a.border, b.border)
                            self.assertEqual(a.alignment, b.alignment)
                            self.assertEqual(a.number_format, b.number_format)
                finally:
                    exported.close()
                    new.close()
            finally:
                rows.close()
                old.close()


if __name__ == '__main__':
    unittest.main()
