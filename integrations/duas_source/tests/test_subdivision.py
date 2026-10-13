"""Failed ranges become durable smaller queries; no live portal requests."""
from collections import Counter
from pathlib import Path
import sqlite3
import sys
import tempfile
from threading import Event
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from data_store import RangeStore, REGIMENES, ranges_for
from runner import run_downloads
from main import unir_archivos_finales
from openpyxl import load_workbook
from tica_dua_bot import Cancelled


def record(day):
    return {'aduana': '001', 'ano': '2026', 'numero': str(day)}


class SubdivisionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.path = self.root / 'duas.sqlite3'
        self.store = RangeStore(self.path)
        self.addCleanup(self.store.close)
        self.cancel = Event()
        self.calls = []

    def run_bot(self, fetch, regimens=('TRANSITO',), refresh=False, start=1, end=15):
        calls = self.calls
        class Browser:
            def __init__(self, *args, **kwargs): pass
            def fetch_range(self, regimen, period, first, last):
                calls.append((regimen, first, last))
                return fetch(first, last)
            def close(self): pass
        with patch.object(self.cancel, 'wait', return_value=False):
            return run_downloads(self.path, {}, regimens, 'Marzo', '2026', 'Rango personalizado',
                                 lambda _: None, self.cancel, start, end, refresh, factory=Browser)

    def test_midpoints_cover_every_month_length_without_gaps(self):
        for period in ('2026-02', '2024-02', '2026-09', '2026-03'):
            for first, last in ranges_for(period, 'Ambos'):
                with self.subTest(period=period, first=first, last=last):
                    parent = self.store.plan(period, ['TRANSITO'], [(first, last)])[0]
                    parts = [self.store.get(key) for key in self.store.split(parent)]
                    boundary = first + (last - first + 1) // 2 - 1
                    self.assertEqual([(p['day_start'], p['day_end']) for p in parts],
                                     [(first, boundary), (boundary + 1, last)])
                    days = [day for p in parts for day in range(p['day_start'], p['day_end'] + 1)]
                    self.assertEqual(days, list(range(first, last + 1)))

    def test_recursive_failures_stop_at_days_and_limit_attempts(self):
        def fetch(first, last):
            if first != last or first == 5:
                raise RuntimeError('heavy range or unavailable day')
            return [record(first)]
        self.assertEqual(self.run_bot(fetch), 1)
        counts = Counter((first, last) for _, first, last in self.calls)
        self.assertEqual(counts[(5, 5)], 3)
        self.assertTrue(all(count == 1 for span, count in counts.items() if span != (5, 5)))
        parent = self.store.plan('2026-03', ['TRANSITO'], [(1, 15)])[0]
        leaves = [self.store.get(key) for key in self.store.leaves([parent])]
        self.assertEqual([(q['day_start'], q['day_end']) for q in leaves], [(day, day) for day in range(1, 16)])
        self.assertEqual(sum(q['status'] == 'succeeded' for q in leaves), 14)
        self.assertFalse(self.store.is_complete(parent))
        # Next run never repeats heavy parents or the fourteen saved days.
        self.calls.clear()
        self.assertEqual(self.run_bot(lambda first, last: [record(first)]), 0)
        self.assertEqual(self.calls, [('TRANSITO', 5, 5)])
        self.assertTrue(self.store.is_complete(parent))

    def test_cancel_resume_reopens_saved_tree(self):
        def fetch(first, last):
            if (first, last) == (1, 15):
                raise RuntimeError('too heavy')
            if first == 8:
                raise Cancelled()
            return [record(first)]
        self.assertEqual(self.run_bot(fetch), 1)
        self.assertEqual(self.calls, [('TRANSITO', 1, 15), ('TRANSITO', 1, 7), ('TRANSITO', 8, 15)])
        self.cancel.clear()
        self.calls.clear()
        self.assertEqual(self.run_bot(lambda first, last: [record(first)]), 0)
        self.assertEqual(self.calls, [('TRANSITO', 8, 15)])

    def test_completed_pieces_export_with_original_classification_and_dedup(self):
        def fetch(first, last):
            if first != last:
                raise RuntimeError('too heavy')
            return [record(first), record(99)]
        self.assertEqual(self.run_bot(fetch, REGIMENES), 0)
        lists, missing = self.store.collect('2026-03', [(1, 15)])
        self.assertFalse(missing)
        for values in lists:
            self.assertEqual(len(values), 16)
            self.assertEqual(len(set(values)), 16)
        paths, missing = unir_archivos_finales('Marzo', '2026', 'Primera parte (1-15)',
                                              lambda _: None, self.root, self.path)
        self.assertFalse(missing)
        self.assertEqual(len(paths), 3)
        for path in paths:
            wb = load_workbook(path, read_only=True)
            try:
                self.assertEqual(wb.sheetnames, ['marzo26'])
                rows = list(wb.active.values)
                self.assertEqual(rows[0], ('NroDUA',))
                self.assertEqual(len(rows), 17)
            finally:
                wb.close()

    def test_missing_child_blocks_export_even_with_old_parent_records(self):
        ids = self.store.plan('2026-03', REGIMENES, [(1, 15)])
        for key in ids:
            self.store.complete(key, [record(99)])
        parent = self.store.plan('2026-03', ['TRANSITO'], [(1, 15)], refresh=True)[0]
        self.store.fail(parent, 'timeout')
        first, last = self.store.split(parent)
        self.store.complete(first, [])
        self.assertEqual(len(self.store.collect('2026-03', [(1, 15)])[1]), 1)
        self.store.complete(last, [record(2)])
        lists, missing = self.store.collect('2026-03', [(1, 15)])
        self.assertFalse(missing)
        self.assertEqual(lists[1], ['001-2026-000002'])

    def test_refresh_invalidates_saved_descendants(self):
        parent = self.store.plan('2026-03', ['TRANSITO'], [(1, 15)])[0]
        left, right = self.store.split(parent)
        children = self.store.split(left) + [right]
        for child in children:
            self.store.complete(child, [record(99)])
        self.assertTrue(self.store.is_complete(parent))
        self.calls.clear()
        self.assertEqual(self.run_bot(lambda first, last: [record(first)], refresh=True), 0)
        self.assertEqual([(a, b) for _, a, b in self.calls], [(1, 3), (4, 7), (8, 15)])
        values = self.store.collect('2026-03', [(1, 15)])[0][1]
        self.assertNotIn('001-2026-000099', values)

    def test_split_reuses_verified_children_and_is_atomic(self):
        parent, child = self.store.plan('2026-03', ['TRANSITO'], [(1, 15), (1, 7)])
        self.store.complete(child, [record(1)])
        self.store.db.execute('''CREATE TRIGGER fail_split BEFORE INSERT ON query_parts
            WHEN (SELECT day_start FROM queries WHERE id=NEW.child_id)=8
            BEGIN SELECT RAISE(ABORT, 'simulated disk failure'); END''')
        with self.assertRaises(sqlite3.IntegrityError):
            self.store.split(parent)
        self.assertEqual(self.store.children(parent), [])
        self.assertEqual(self.store.db.execute('SELECT COUNT(*) FROM queries').fetchone()[0], 2)
        self.store.db.execute('DROP TRIGGER fail_split')
        parts = self.store.split(parent)
        self.assertEqual(parts[0], child)
        self.assertEqual(self.store.split(parent), parts)
        self.assertEqual(self.store.get(child)['status'], 'succeeded')

    def test_browser_launch_failure_does_not_split_every_day(self):
        class Broken:
            def __init__(self, *args, **kwargs):
                raise RuntimeError('driver unavailable')
        with patch.object(self.cancel, 'wait', return_value=False):
            self.assertEqual(run_downloads(self.path, {}, ['TRANSITO'], 'Marzo', '2026',
                             'Primera parte (1-15)', lambda _: None, self.cancel, factory=Broken), 1)
        self.assertEqual(self.store.db.execute('SELECT COUNT(*) FROM queries').fetchone()[0], 1)
        self.assertEqual(self.store.db.execute('SELECT attempts FROM queries').fetchone()[0], 3)

    def test_refresh_also_finishes_older_custom_ranges(self):
        ids = self.store.plan('2026-03', ['TRANSITO'], [(1, 15), (1, 7), (8, 15)])
        for key in ids:
            self.store.complete(key, [record(99)])
        self.assertEqual(self.run_bot(lambda first, last: [record(first)], refresh=True), 0)
        self.assertEqual(len(self.calls), 3)
        values, missing = self.store.collect('2026-03', [(1, 15)])
        self.assertNotIn('001-2026-000099', values[1])
        self.assertFalse(any(message.startswith('TRANSITO:') for message in missing))

    def test_opening_previous_schema_keeps_saved_queries_and_records(self):
        key = self.store.plan('2026-03', ['TRANSITO'], [(1, 15)])[0]
        self.store.complete(key, [record(1)])
        self.store.db.execute('DROP TABLE query_parts')
        reopened = RangeStore(self.path)
        try:
            self.assertTrue(reopened.is_complete(key))
            self.assertEqual(reopened.collect('2026-03', [(1, 15)])[0][1], ['001-2026-000001'])
            self.assertEqual(reopened.db.execute('SELECT COUNT(*) FROM query_parts').fetchone()[0], 0)
        finally:
            reopened.close()


if __name__ == '__main__':
    unittest.main()
