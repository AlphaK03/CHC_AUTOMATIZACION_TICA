"""Extrae dos DUAs confirmados de la base original en modo solo lectura."""
import argparse
import hashlib
import json
from pathlib import Path
import sqlite3

from openpyxl import Workbook


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    folder = args.source.resolve() / 'data' / 'queues'
    session = json.loads((args.source / 'config' / 'session.json').read_text(encoding='utf-8'))
    catalog = sqlite3.connect((folder / 'catalog.sqlite3').as_uri() + '?mode=ro', uri=True)
    catalog.row_factory = sqlite3.Row
    try:
        table = catalog.execute('SELECT * FROM saved_tables WHERE id=? AND deleted=0', (session['active_table_id'],)).fetchone()
        if table is None:
            raise ValueError('No existe la tabla activa original.')
        table = dict(table)
    finally:
        catalog.close()
    path = (folder / table['filename']).resolve()
    if path.parent != folder:
        raise ValueError('Ruta de tabla fuera del catálogo original.')
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    db = sqlite3.connect(path.as_uri() + '?mode=ro', uri=True)
    db.row_factory = sqlite3.Row
    try:
        db.execute('BEGIN')
        rows = [dict(row) for row in db.execute("SELECT dua,result FROM jobs WHERE status='succeeded' ORDER BY rowid LIMIT 2")]
        counts = dict(db.execute('SELECT status,COUNT(*) FROM jobs GROUP BY status').fetchall())
    finally:
        db.close()
    if len(rows) != 2:
        raise ValueError('Se necesitan dos resultados confirmados para la comparación.')
    args.output.mkdir(parents=True, exist_ok=True)
    expected = {row['dua']: json.loads(row['result']) for row in rows}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'muestra_original'
    sheet.append(['NroDUA'])
    for dua in expected:
        sheet.append([dua])
        sheet.cell(sheet.max_row, 1).number_format = '@'
    workbook.save(args.output / 'muestra_original.xlsx')
    workbook.close()
    after = hashlib.sha256(path.read_bytes()).hexdigest()
    reference = {'source_database': str(path), 'source_table_name': table['name'],
                 'original_counts': counts, 'source_sha256_before': before,
                 'source_sha256_after': after, 'expected': expected}
    (args.output / 'referencia.json').write_text(json.dumps(reference, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'sample_rows': len(expected), 'original_counts': counts,
                      'source_unchanged': before == after, 'sample_path': str(args.output / 'muestra_original.xlsx')}, indent=2))


if __name__ == '__main__':
    main()
